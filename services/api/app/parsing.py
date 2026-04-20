from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.models import DocumentRecord, DocumentType

MONTH_PATTERN = (
    r"(?:January|February|March|April|May|June|July|August|September|October|"
    r"November|December)\s+\d{1,2},\s+\d{4}"
)
ENTITY_PATTERN = re.compile(
    r"\b([A-Z][A-Za-z0-9&.,'()\- ]{2,80}?"
    r"(?:Inc\.|Incorporated|LLC|Ltd\.|Limited|LP|L\.P\.|Corp\.|Corporation|"
    r"Company|Co\.|PLC|Holdings|Partners|Partnership))\b"
)


@dataclass(slots=True)
class ParsedPage:
    page_number: int
    section_heading: str
    text: str


@dataclass(slots=True)
class ParsedClause:
    page_number: int
    section_heading: str
    text: str


@dataclass(slots=True)
class ParsedDocument:
    document: DocumentRecord
    pages: list[ParsedPage]
    clauses: list[ParsedClause]
    text: str


def parse_document(document_id: str, filename: str, file_path: Path) -> ParsedDocument:
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        pages = parse_pdf(file_path)
        clauses = extract_pdf_clauses(pages)
    elif suffix == ".docx":
        pages, clauses = parse_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    full_text = "\n\n".join(page.text for page in pages).strip()
    doc_type = classify_document(filename, full_text)

    document = DocumentRecord(
        id=document_id,
        name=filename,
        doc_type=doc_type,
        counterparty=extract_counterparty(filename, full_text),
        effective_date=extract_effective_date(full_text),
        expiry_date=extract_expiry_date(full_text),
        renewal_notice_days=extract_renewal_notice_days(full_text),
        auto_renews=detect_auto_renewal(full_text),
        governing_law=extract_governing_law(full_text),
    )

    return ParsedDocument(document=document, pages=pages, clauses=clauses, text=full_text)


def parse_pdf(file_path: Path) -> list[ParsedPage]:
    reader = PdfReader(str(file_path))
    pages: list[ParsedPage] = []

    for index, page in enumerate(reader.pages, start=1):
        text = normalize_whitespace(page.extract_text() or "")
        if not text:
            continue

        pages.append(
            ParsedPage(
                page_number=index,
                section_heading=guess_heading(text),
                text=text,
            )
        )

    return pages or [ParsedPage(page_number=1, section_heading="Document text", text="")]


def parse_docx(file_path: Path) -> tuple[list[ParsedPage], list[ParsedClause]]:
    document = DocxDocument(str(file_path))
    page_number = 1
    current_heading = "Document text"
    buckets: dict[int, list[str]] = {1: []}
    headings: dict[int, str] = {1: current_heading}
    clauses: list[ParsedClause] = []

    for paragraph in document.paragraphs:
        text = normalize_whitespace(paragraph.text)
        style_name = paragraph.style.name if paragraph.style else ""

        if style_name.lower().startswith("heading") and text:
            current_heading = text
            headings[page_number] = current_heading

        if text:
            buckets.setdefault(page_number, []).append(text)
            if not style_name.lower().startswith("heading"):
                clauses.append(
                    ParsedClause(
                        page_number=page_number,
                        section_heading=current_heading,
                        text=text,
                    )
                )

        if paragraph_has_page_break(paragraph):
            page_number += 1
            buckets.setdefault(page_number, [])
            headings.setdefault(page_number, current_heading)

    for table in document.tables:
        rows = []
        for row in table.rows:
            row_text = " | ".join(
                normalize_whitespace(cell.text) for cell in row.cells if cell.text.strip()
            )
            if row_text:
                rows.append(row_text)
                clauses.append(
                    ParsedClause(
                        page_number=page_number,
                        section_heading=current_heading,
                        text=row_text,
                    )
                )
        if rows:
            buckets.setdefault(page_number, []).append("\n".join(rows))

    pages = [
        ParsedPage(
            page_number=number,
            section_heading=headings.get(number, "Document text"),
            text="\n\n".join(content).strip(),
        )
        for number, content in sorted(buckets.items())
        if "\n\n".join(content).strip()
    ]

    fallback_pages = pages or [ParsedPage(page_number=1, section_heading="Document text", text="")]
    fallback_clauses = clauses or extract_pdf_clauses(fallback_pages)
    return fallback_pages, fallback_clauses


def extract_pdf_clauses(pages: list[ParsedPage]) -> list[ParsedClause]:
    clauses: list[ParsedClause] = []
    for page in pages:
        for segment in split_into_clause_units(page.text):
            clauses.append(
                ParsedClause(
                    page_number=page.page_number,
                    section_heading=page.section_heading,
                    text=segment,
                )
            )
    return clauses


def paragraph_has_page_break(paragraph: object) -> bool:
    xml = getattr(getattr(paragraph, "_p", None), "xml", "")
    return 'w:type="page"' in xml or "lastRenderedPageBreak" in xml


def split_into_clause_units(value: str) -> list[str]:
    normalized = value.replace("\r", "\n")
    parts = [
        normalize_whitespace(part)
        for part in re.split(r"\n{2,}", normalized)
        if normalize_whitespace(part)
    ]
    clauses = [part[:650] for part in parts if len(part) >= 35]
    if clauses:
        return clauses[:8]
    fallback = normalize_whitespace(value)
    return [fallback[:650]] if fallback else []


def classify_document(filename: str, text: str) -> DocumentType:
    haystack = f"{filename}\n{text[:5000]}".lower()

    if "lease" in haystack:
        return DocumentType.LEASE
    if "nda" in haystack or "non-disclosure" in haystack or "confidentiality agreement" in haystack:
        return DocumentType.NDA
    if "employment" in haystack or "offer letter" in haystack:
        return DocumentType.EMPLOYMENT
    if "amendment" in haystack or "amended and restated" in haystack:
        return DocumentType.AMENDMENT
    if "software as a service" in haystack or "saas" in haystack:
        return DocumentType.SAAS
    if any(
        token in haystack
        for token in [
            "vendor",
            "supplier",
            "service provider",
            "hosting",
            "distribution",
            "distributor",
            "data processing addendum",
            "dpa",
            "partnership",
        ]
    ):
        return DocumentType.VENDOR_AGREEMENT
    if any(
        token in haystack
        for token in [
            "customer",
            "client",
            "master services agreement",
            "msa",
            "order form",
            "statement of work",
        ]
    ):
        return DocumentType.CUSTOMER_AGREEMENT

    return DocumentType.OTHER


def extract_counterparty(filename: str, text: str) -> str | None:
    preview = text[:5000]
    match = ENTITY_PATTERN.search(preview)
    if match:
        return normalize_whitespace(match.group(1))

    stem = Path(filename).stem
    simplified = re.sub(
        r"(?i)\b(msa|agreement|contract|addendum|dpa|services|service|master|lease|nda|amendment)\b",
        "",
        stem,
    )
    simplified = normalize_whitespace(simplified.replace("-", " ").replace("_", " "))
    return simplified or None


def extract_effective_date(text: str) -> str | None:
    patterns = [
        re.compile(
            rf"(?:effective date|effective as of|dated as of|made as of)\s*[:\-]?\s*({MONTH_PATTERN})",
            re.IGNORECASE,
        ),
        re.compile(
            r"(?:effective date|effective as of|dated as of|made as of)\s*[:\-]?\s*(\d{4}-\d{2}-\d{2})",
            re.IGNORECASE,
        ),
    ]
    return first_match_group(patterns, text)


def extract_expiry_date(text: str) -> str | None:
    patterns = [
        re.compile(
            rf"(?:expires on|expiration date|term ends on|expiry date)\s*[:\-]?\s*({MONTH_PATTERN})",
            re.IGNORECASE,
        ),
        re.compile(
            r"(?:expires on|expiration date|term ends on|expiry date)\s*[:\-]?\s*(\d{4}-\d{2}-\d{2})",
            re.IGNORECASE,
        ),
    ]
    return first_match_group(patterns, text)


def extract_renewal_notice_days(text: str) -> int | None:
    pattern = re.compile(
        r"(\d{1,3})\s+days?.{0,80}(?:before|prior).{0,80}(?:renew|term)",
        re.IGNORECASE | re.DOTALL,
    )
    match = pattern.search(text)
    if not match:
        return None

    try:
        return int(match.group(1))
    except ValueError:
        return None


def detect_auto_renewal(text: str) -> bool:
    return bool(
        re.search(
            r"automatic(?:ally)? renew|renew automatically|additional one[- ]year term|successive renewal term",
            text,
            re.IGNORECASE,
        )
    )


def extract_governing_law(text: str) -> str | None:
    patterns = [
        re.compile(
            r"governed by the laws of ([A-Za-z ,()]+?)(?:\.|;|\n)",
            re.IGNORECASE,
        ),
        re.compile(
            r"laws of the Province of ([A-Za-z ]+?)(?:\.|;|\n)",
            re.IGNORECASE,
        ),
    ]
    return first_match_group(patterns, text)


def guess_heading(text: str) -> str:
    for line in text.splitlines():
        candidate = normalize_whitespace(line)
        if candidate and len(candidate) <= 120:
            return candidate
    return "Matched clause"


def first_match_group(patterns: list[re.Pattern[str]], text: str) -> str | None:
    for pattern in patterns:
        match = pattern.search(text)
        if match:
            return normalize_whitespace(match.group(1))
    return None


def normalize_whitespace(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()
