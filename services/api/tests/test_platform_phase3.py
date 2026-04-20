from __future__ import annotations

import importlib
import json
import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient


FIXTURES_DIR = Path(__file__).parent / "fixtures" / "parse"


def _reload_app_modules() -> None:
    for name in list(sys.modules):
        if name == "app" or name.startswith("app."):
            sys.modules.pop(name, None)


def _patch_local_storage(
    *,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    storage,
    repository,
    object_storage,
) -> None:
    data_dir = tmp_path / "data"
    uploads_dir = tmp_path / "uploads"
    db_path = data_dir / "skua.db"

    monkeypatch.setattr(storage, "DATA_DIR", data_dir)
    monkeypatch.setattr(storage, "UPLOADS_DIR", uploads_dir)
    monkeypatch.setattr(storage, "DB_PATH", db_path)
    monkeypatch.setattr(repository, "UPLOADS_DIR", uploads_dir)
    monkeypatch.setattr(object_storage, "LOCAL_SOURCE_DIR", uploads_dir / "source")
    monkeypatch.setattr(object_storage, "LOCAL_ARTIFACT_DIR", uploads_dir / "artifacts")


@pytest.fixture()
def platform_client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setenv("SKUA_DATABASE_URL", f"sqlite+pysqlite:///{tmp_path / 'platform.db'}")
    monkeypatch.setenv("SKUA_QUEUE_BACKEND", "sqlite-only")
    monkeypatch.setenv("SKUA_DEV_USER_EMAIL", "founder@test.local")
    monkeypatch.setenv("SKUA_DEV_USER_PASSWORD", "changeme123")

    _reload_app_modules()
    main = importlib.import_module("app.main")
    repository = importlib.import_module("app.repository")
    storage = importlib.import_module("app.storage")
    object_storage = importlib.import_module("app.object_storage")

    _patch_local_storage(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        storage=storage,
        repository=repository,
        object_storage=object_storage,
    )

    with TestClient(main.app) as client:
        yield client


def _register_user(client: TestClient) -> tuple[dict[str, str], str]:
    response = client.post(
        "/api/v1/auth/register",
        json={
            "email": "phase3@example.com",
            "password": "changeme123",
            "full_name": "Phase Three",
        },
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    return {"Authorization": f"Bearer {payload['access_token']}"}, payload["user"]["workspace_ids"][0]


def _build_docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("Confidentiality", level=1)
    document.add_paragraph("1. Use Restriction. Recipient may use confidential information only to evaluate the transaction.")
    document.add_paragraph("Recipient will protect the information using reasonable safeguards and return it on request.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Party"
    table.cell(0, 1).text = "Notice Days"
    table.cell(1, 0).text = "Recipient"
    table.cell(1, 1).text = "10"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_simple_pdf_bytes(lines: list[str]) -> bytes:
    objects: list[bytes] = []
    content_lines = ["BT", "/F1 12 Tf", "72 720 Td"]
    first = True
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        if first:
            content_lines.append(f"({escaped}) Tj")
            first = False
        else:
            content_lines.append("0 -18 Td")
            content_lines.append(f"({escaped}) Tj")
    content_lines.append("ET")
    stream = "\n".join(content_lines).encode("utf-8")

    objects.append(b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n")
    objects.append(b"2 0 obj\n<< /Type /Pages /Count 1 /Kids [3 0 R] >>\nendobj\n")
    objects.append(
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
    )
    objects.append(
        f"4 0 obj\n<< /Length {len(stream)} >>\nstream\n".encode("utf-8")
        + stream
        + b"\nendstream\nendobj\n"
    )
    objects.append(b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n")

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objects:
        offsets.append(len(output))
        output.extend(obj)

    xref_start = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("utf-8"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("utf-8"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_start}\n%%EOF"
        ).encode("utf-8")
    )
    return bytes(output)


def test_platform_docx_upload_matches_golden_fixture(platform_client: TestClient) -> None:
    headers, workspace_id = _register_user(platform_client)
    response = platform_client.post(
        "/api/v1/platform/documents/upload",
        headers=headers,
        data={"workspace_id": workspace_id, "source_kind": "word_document"},
        files={
            "files": (
                "confidentiality.docx",
                _build_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200, response.text
    payload = response.json()[0]
    assert payload["parse_status"] == "completed"
    assert payload["document_version"]["status"] == "indexed"

    detail_response = platform_client.get(
        f"/api/v1/platform/document-versions/{payload['document_version']['id']}",
        headers=headers,
    )
    assert detail_response.status_code == 200, detail_response.text
    detail = detail_response.json()

    expected = json.loads((FIXTURES_DIR / "docx_contract.expected.json").read_text("utf-8"))
    assert detail["parser_name"] == expected["parser_name"]
    assert [segment["segment_type"] for segment in detail["segments"]] == expected["segment_types"]
    assert [segment["page_number"] for segment in detail["segments"]] == expected["page_numbers"]


def test_platform_selection_ingest_search_and_anchor_relocation(platform_client: TestClient) -> None:
    headers, workspace_id = _register_user(platform_client)
    response = platform_client.post(
        "/api/v1/platform/documents/selection",
        headers=headers,
        json={
            "workspace_id": workspace_id,
            "document_name": "selection.txt",
            "selection_text": (
                "ASSIGNMENT\n\n"
                "1. Assignment. Neither party may assign this Agreement without prior written consent.\n\n"
                "2. Change of Control. A change of control is deemed an assignment."
            ),
        },
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    document_version_id = payload["document_version"]["id"]

    search_response = platform_client.post(
        f"/api/v1/platform/document-versions/{document_version_id}/search",
        headers=headers,
        json={"query": "assignment consent", "limit": 3, "segment_types": ["clause", "paragraph"]},
    )
    assert search_response.status_code == 200, search_response.text
    search_results = search_response.json()
    assert search_results
    assert "assign" in search_results[0]["text"].lower()

    relocate_response = platform_client.post(
        "/api/v1/platform/anchors/relocate",
        headers=headers,
        json={
            "document_version_id": document_version_id,
            "anchor": search_results[0]["anchor"],
        },
    )
    assert relocate_response.status_code == 200, relocate_response.text
    relocation = relocate_response.json()
    assert relocation["strategy"] in {"exact_match", "quote_hash_match"}
    assert relocation["ordinal"] == search_results[0]["ordinal"]


def test_pdf_parser_matches_golden_fixture(tmp_path: Path) -> None:
    platform_parsing = importlib.import_module("app.platform_parsing")
    pdf_path = tmp_path / "contract.pdf"
    pdf_path.write_bytes(
        _build_simple_pdf_bytes(
            [
                "CONFIDENTIALITY AGREEMENT",
                "",
                "1. Term. The term is one year and renews automatically unless notice is given.",
            ]
        )
    )
    parse_result = platform_parsing.parse_pdf_document(
        document_version_id="pdv-test",
        file_path=pdf_path,
    )

    expected = json.loads((FIXTURES_DIR / "pdf_contract.expected.json").read_text("utf-8"))
    assert parse_result.parser_name == expected["parser_name"]
    assert [segment.segment_type for segment in parse_result.segments] == expected["segment_types"]
    assert [segment.page_number for segment in parse_result.segments] == expected["page_numbers"]
