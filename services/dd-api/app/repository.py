from __future__ import annotations

import hashlib
import json
import re
from csv import writer
from datetime import UTC, datetime
from functools import lru_cache
from io import BytesIO, StringIO
from pathlib import Path
from uuid import uuid4

from docx import Document as DocxDocument
from docx.shared import Pt
from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl.styles import Alignment, Font, PatternFill
import yaml

from app.analyzer import (
    build_highlights,
    generate_issues,
    generate_output,
    generate_review_suggestions,
)
from app.models import (
    AnswerType,
    AnchorReconciliationStatus,
    AskRunCreateRequest,
    AskRunRecord,
    AskSourceToggles,
    AuditEventRecord,
    CanonicalUploadDocumentsResponse,
    CitationRecord,
    DocumentRecord,
    DocumentAnchor,
    DocumentType,
    DocumentVersionDetail,
    DocumentVersionRecord,
    DraftLibraryMatch,
    DraftMode,
    DraftRunCreateRequest,
    DraftRunRecord,
    GeneratedOutput,
    IssueRecord,
    IssueStatus,
    JobRecord,
    JobStatus,
    LibraryItemRecord,
    LibrarySearchRequest,
    PlaybookSavedNoteRecord,
    PlaybookRecord,
    ProjectRecord,
    ProjectCreateRequest,
    QueryRunCell,
    QueryRunCreateRequest,
    QueryRunRecord,
    QueryRunRow,
    DdReportDocumentSummary,
    DdReportEventRecord,
    DdReportEventSummary,
    DdReportRecord,
    DdReportUpdateRequest,
    ProposedRedline,
    ReviewAudience,
    ReviewSuggestionApplyRequest,
    ReviewSuggestionDismissRequest,
    ReviewSuggestionEvent,
    ReviewSuggestionMarkReviewedRequest,
    ReviewSuggestionSaveToPlaybookRequest,
    ReviewRunExportRecord,
    ReviewRunCreateRequest,
    ReviewRunRecord,
    ReviewRunSummary,
    ReviewSuggestionRecord,
    ReviewType,
    SeverityLevel,
    StandardsRunCreateRequest,
    StandardsRunRecord,
    StandardsTemplateRecord,
    StandardsWeakClause,
    SuggestionStatus,
    UploadDocumentsResponse,
    WorkflowRunCreateRequest,
    WorkflowArtifactVariantTemplate,
    WorkflowRunEventRecord,
    WorkflowRunEventSummary,
    WorkflowRunRecord,
    WorkflowRunRerunRequest,
    WorkflowRunUpdateRequest,
    WorkflowWorkbookSheetTemplate,
    WorkflowMemoSectionTemplate,
    WorkflowTemplateRecord,
    WorkspaceDetail,
    WorkspaceCreateRequest,
    WorkspaceSummary,
)
from app.parsing import ParsedClause, ParsedPage, parse_document
from app.seed import SEED_OUTPUTS, SEED_WORKSPACES
from app.storage import ROOT_DIR, UPLOADS_DIR, ensure_storage, get_connection

PLAYBOOK_DIR = ROOT_DIR / "packages" / "playbooks"
WORKFLOW_DIR = ROOT_DIR / "packages" / "workflows"
STANDARDS_DIR = ROOT_DIR / "packages" / "standards"


@lru_cache
def load_playbooks() -> list[PlaybookRecord]:
    playbooks: list[PlaybookRecord] = []

    for playbook_path in sorted(PLAYBOOK_DIR.glob("*.yaml")):
        with playbook_path.open("r", encoding="utf-8") as handle:
            payload = yaml.safe_load(handle)

        playbooks.append(PlaybookRecord.model_validate(payload))

    return playbooks


@lru_cache
def load_workflow_templates() -> list[WorkflowTemplateRecord]:
    templates: list[WorkflowTemplateRecord] = []

    for workflow_path in sorted(WORKFLOW_DIR.glob("*.yaml")):
        with workflow_path.open("r", encoding="utf-8") as handle:
            payload = yaml.safe_load(handle)

        templates.append(normalize_workflow_template(WorkflowTemplateRecord.model_validate(payload)))

    return templates


@lru_cache
def load_standards_templates() -> list[StandardsTemplateRecord]:
    templates: list[StandardsTemplateRecord] = []

    for standards_path in sorted(STANDARDS_DIR.glob("*.yaml")):
        with standards_path.open("r", encoding="utf-8") as handle:
            payload = yaml.safe_load(handle)

        templates.append(StandardsTemplateRecord.model_validate(payload))

    return templates


def normalize_workflow_template(template: WorkflowTemplateRecord) -> WorkflowTemplateRecord:
    artifact_variants = template.artifact_variants or [
        WorkflowArtifactVariantTemplate(
            id="default",
            name="Default",
            description="Default report layout for this workflow.",
        )
    ]
    workbook_sheets = template.workbook_sheets or [
        WorkflowWorkbookSheetTemplate(id="workflow-results", title="Workflow Results", kind="results"),
        WorkflowWorkbookSheetTemplate(id="citations", title="Citations", kind="citations"),
        WorkflowWorkbookSheetTemplate(id="exceptions", title="Exceptions", kind="exceptions"),
        WorkflowWorkbookSheetTemplate(
            id="document-summaries",
            title="Document Summaries",
            kind="document_summaries",
        ),
        WorkflowWorkbookSheetTemplate(id="history", title="History", kind="history"),
    ]
    memo_sections = template.memo_sections or [
        WorkflowMemoSectionTemplate(heading="Executive summary", kind="overview"),
        WorkflowMemoSectionTemplate(
            heading="Document summaries",
            kind="document_summaries",
        ),
        WorkflowMemoSectionTemplate(heading="Exceptions list", kind="exceptions"),
    ]
    return template.model_copy(
        update={
            "default_artifact_variant_id": template.default_artifact_variant_id or artifact_variants[0].id,
            "artifact_variants": artifact_variants,
            "workbook_sheets": workbook_sheets,
            "memo_sections": memo_sections,
        }
    )


def resolve_workflow_artifact_variant(
    template: WorkflowTemplateRecord,
    artifact_variant_id: str | None,
) -> tuple[WorkflowArtifactVariantTemplate | None, list[WorkflowWorkbookSheetTemplate], list[WorkflowMemoSectionTemplate]]:
    variant = next(
        (
            candidate
            for candidate in template.artifact_variants
            if candidate.id == (artifact_variant_id or template.default_artifact_variant_id)
        ),
        template.artifact_variants[0] if template.artifact_variants else None,
    )
    workbook_sheets = variant.workbook_sheets if variant and variant.workbook_sheets else template.workbook_sheets
    memo_sections = variant.memo_sections if variant and variant.memo_sections else template.memo_sections
    return variant, workbook_sheets, memo_sections


def init_repository() -> None:
    ensure_storage()
    with get_connection() as connection:
        orphan_rows = connection.execute(
            """
            SELECT w.id, w.name, w.stage, w.created_at, w.last_updated
            FROM workspaces w
            LEFT JOIN projects p ON p.workspace_id = w.id
            WHERE p.id IS NULL
            """
        ).fetchall()
        for row in orphan_rows:
            connection.execute(
                """
                INSERT INTO projects (id, workspace_id, name, stage, created_at, last_updated)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    f"proj-{uuid4().hex[:12]}",
                    row["id"],
                    row["name"],
                    row["stage"],
                    row["created_at"],
                    row["last_updated"],
                ),
            )
        orphan_documents = connection.execute(
            """
            SELECT d.id, d.workspace_id, d.name, d.file_path, d.file_hash, d.extracted_text, d.created_at
            FROM documents d
            LEFT JOIN document_versions dv ON dv.source_document_id = d.id
            WHERE dv.id IS NULL
            """
        ).fetchall()
        for row in orphan_documents:
            project_row = connection.execute(
                "SELECT id FROM projects WHERE workspace_id = ?",
                (row["workspace_id"],),
            ).fetchone()
            if project_row is None:
                continue

            document_version_id = f"dv-{uuid4().hex[:12]}"
            connection.execute(
                """
                INSERT INTO document_versions (
                    id, project_id, source_document_id, name, source_type, file_path, file_hash,
                    extracted_text, status, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    document_version_id,
                    project_row["id"],
                    row["id"],
                    row["name"],
                    "upload",
                    row["file_path"],
                    row["file_hash"],
                    row["extracted_text"],
                    "ingested",
                    row["created_at"],
                ),
            )
            page_rows = connection.execute(
                """
                SELECT page_number, section_heading, text
                FROM document_pages
                WHERE document_id = ?
                ORDER BY page_number ASC
                """,
                (row["id"],),
            ).fetchall()
            insert_document_anchors(
                connection=connection,
                document_version_id=document_version_id,
                pages=[
                    ParsedPage(
                        page_number=page_row["page_number"],
                        section_heading=page_row["section_heading"],
                        text=page_row["text"],
                    )
                    for page_row in page_rows
                ],
            )
        orphan_library_versions = connection.execute(
            """
            SELECT dv.id, dv.project_id, dv.name, d.doc_type, d.governing_law, d.counterparty
            FROM document_versions dv
            JOIN documents d ON d.id = dv.source_document_id
            LEFT JOIN library_items li ON li.document_version_id = dv.id
            WHERE li.id IS NULL
            """
        ).fetchall()
        for row in orphan_library_versions:
            anchor_rows = connection.execute(
                """
                SELECT id, document_version_id, page_number, quote, metadata_json
                FROM document_anchors
                WHERE document_version_id = ?
                ORDER BY created_at ASC
                """,
                (row["id"],),
            ).fetchall()
            insert_library_items(
                connection=connection,
                project_id=row["project_id"],
                document_version_id=row["id"],
                document_name=row["name"],
                doc_type=DocumentType(row["doc_type"]),
                governing_law=row["governing_law"],
                counterparty=row["counterparty"],
                anchor_rows=anchor_rows,
            )


def list_workspaces() -> list[WorkspaceSummary]:
    init_repository()
    workspaces = [detail.workspace for detail in SEED_WORKSPACES.values()]

    with get_connection() as connection:
        rows = connection.execute(
            "SELECT id, name, stage, playbook_names_json, last_updated FROM workspaces ORDER BY last_updated DESC"
        ).fetchall()

        for row in rows:
            workspaces.append(build_workspace_summary(connection, row))

    return workspaces


def list_projects() -> list[ProjectRecord]:
    init_repository()

    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, workspace_id, name, stage, created_at, last_updated
            FROM projects
            ORDER BY last_updated DESC
            """
        ).fetchall()
        return [build_project_record(connection, row) for row in rows]


def create_project(payload: ProjectCreateRequest) -> ProjectRecord:
    workspace = create_workspace(payload.name)
    project = get_project_by_workspace_id(workspace.id)
    if project is None:
        raise ValueError("Project creation failed.")
    return project


def create_workspace(name: str) -> WorkspaceSummary:
    init_repository()
    workspace_id = build_workspace_id(name)
    timestamp = now_timestamp()
    playbook_names = [playbook.name for playbook in load_playbooks()]
    project_id = f"proj-{uuid4().hex[:12]}"

    with get_connection() as connection:
        connection.execute(
            """
            INSERT INTO workspaces (id, name, stage, playbook_names_json, created_at, last_updated)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                workspace_id,
                name.strip(),
                "Uploaded documents pending review",
                json.dumps(playbook_names),
                timestamp,
                timestamp,
            ),
        )
        connection.execute(
            """
            INSERT INTO projects (id, workspace_id, name, stage, created_at, last_updated)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                project_id,
                workspace_id,
                name.strip(),
                "Uploaded documents pending review",
                timestamp,
                timestamp,
            ),
        )
        insert_audit_event(
            connection=connection,
            event_type="project.created",
            actor_surface="system",
            project_id=project_id,
            document_version_id=None,
            review_run_id=None,
            suggestion_id=None,
            payload={"workspace_id": workspace_id, "name": name.strip()},
        )
        row = connection.execute(
            "SELECT id, name, stage, playbook_names_json, last_updated FROM workspaces WHERE id = ?",
            (workspace_id,),
        ).fetchone()

        return build_workspace_summary(connection, row)


def get_project(project_id: str) -> ProjectRecord | None:
    init_repository()

    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, workspace_id, name, stage, created_at, last_updated
            FROM projects
            WHERE id = ?
            """,
            (project_id,),
        ).fetchone()
        if row is None:
            return None
        return build_project_record(connection, row)


def get_project_by_workspace_id(workspace_id: str) -> ProjectRecord | None:
    init_repository()

    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, workspace_id, name, stage, created_at, last_updated
            FROM projects
            WHERE workspace_id = ?
            """,
            (workspace_id,),
        ).fetchone()
        if row is None:
            return None
        return build_project_record(connection, row)


def get_workspace(workspace_id: str) -> WorkspaceDetail | None:
    init_repository()

    if workspace_id in SEED_WORKSPACES:
        return SEED_WORKSPACES.get(workspace_id)

    with get_connection() as connection:
        row = connection.execute(
            "SELECT id, name, stage, playbook_names_json, last_updated FROM workspaces WHERE id = ?",
            (workspace_id,),
        ).fetchone()
        if row is None:
            return None

        return build_workspace_detail(connection, row)


def get_first_pass_output(workspace_id: str) -> GeneratedOutput | None:
    if workspace_id in SEED_OUTPUTS:
        return SEED_OUTPUTS.get(workspace_id)

    workspace = get_workspace(workspace_id)
    if workspace is None:
        return None

    return generate_output(workspace)


def create_draft_run(payload: DraftRunCreateRequest) -> DraftRunRecord:
    init_repository()
    draft_run_id = f"drf-{uuid4().hex[:12]}"
    created_at = now_timestamp()

    with get_connection() as connection:
        job_id = create_job_record(
            connection=connection,
            job_type="draft_run",
            status=JobStatus.QUEUED,
            project_id=payload.project_id,
            document_version_id=payload.document_version_id,
            export_id=None,
            metadata={"draft_run_id": draft_run_id, "mode": str(payload.mode)},
        )
        connection.execute(
            """
            INSERT INTO draft_runs (
                id, project_id, document_version_id, job_id, mode, query, instruction,
                selection_text, selection_anchor_json, status, generated_text,
                citations_json, library_matches_json, created_at, completed_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                draft_run_id,
                payload.project_id,
                payload.document_version_id,
                job_id,
                payload.mode,
                payload.query.strip() if payload.query else None,
                payload.instruction.strip() if payload.instruction else None,
                payload.selection_text.strip() if payload.selection_text else None,
                json.dumps(payload.selection_anchor.model_dump() if payload.selection_anchor else None),
                JobStatus.QUEUED,
                "",
                json.dumps([]),
                json.dumps([]),
                created_at,
                None,
            ),
        )
        insert_audit_event(
            connection=connection,
            event_type="draft_run.queued",
            actor_surface="word_addin",
            project_id=payload.project_id,
            document_version_id=payload.document_version_id,
            review_run_id=None,
            suggestion_id=None,
            payload={"draft_run_id": draft_run_id, "job_id": job_id, "mode": str(payload.mode)},
        )
        draft_row = connection.execute(
            """
            SELECT id, project_id, document_version_id, job_id, mode, query, instruction,
                   status, generated_text, citations_json, library_matches_json,
                   created_at, completed_at
            FROM draft_runs
            WHERE id = ?
            """,
            (draft_run_id,),
        ).fetchone()
        return build_draft_run_record(draft_row)


def create_ask_run(payload: AskRunCreateRequest) -> AskRunRecord:
    init_repository()
    ask_run_id = f"ask-{uuid4().hex[:12]}"
    created_at = now_timestamp()

    with get_connection() as connection:
        job_id = create_job_record(
            connection=connection,
            job_type="ask_run",
            status=JobStatus.QUEUED,
            project_id=payload.project_id,
            document_version_id=payload.document_version_id,
            export_id=None,
            metadata={"ask_run_id": ask_run_id, "answer_type": str(payload.answer_type)},
        )
        connection.execute(
            """
            INSERT INTO ask_runs (
                id, project_id, document_version_id, job_id, selection_anchor_id,
                selection_text, selection_anchor_json, question, answer_type,
                source_toggles_json, status, answer_markdown, citations_json,
                created_at, completed_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                ask_run_id,
                payload.project_id,
                payload.document_version_id,
                job_id,
                payload.selection_anchor_id,
                payload.selection_text.strip() if payload.selection_text else None,
                json.dumps(payload.selection_anchor.model_dump() if payload.selection_anchor else None),
                payload.question.strip(),
                payload.answer_type,
                json.dumps(payload.source_toggles.model_dump()),
                JobStatus.QUEUED,
                "",
                json.dumps([]),
                created_at,
                None,
            ),
        )
        insert_audit_event(
            connection=connection,
            event_type="ask_run.queued",
            actor_surface="word_addin",
            project_id=payload.project_id,
            document_version_id=payload.document_version_id,
            review_run_id=None,
            suggestion_id=None,
            payload={"ask_run_id": ask_run_id, "job_id": job_id, "answer_type": str(payload.answer_type)},
        )
        ask_row = connection.execute(
            """
            SELECT id, project_id, document_version_id, job_id, question, source_toggles_json,
                   status, answer_type, answer_markdown, citations_json, created_at, completed_at
            FROM ask_runs
            WHERE id = ?
            """,
            (ask_run_id,),
        ).fetchone()
        return build_ask_run_record(ask_row)


def get_draft_run(draft_run_id: str) -> DraftRunRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, document_version_id, job_id, mode, query, instruction,
                   status, generated_text, citations_json, library_matches_json,
                   created_at, completed_at
            FROM draft_runs
            WHERE id = ?
            """,
            (draft_run_id,),
        ).fetchone()
        if row is None:
            return None
        return build_draft_run_record(row)


def get_ask_run(ask_run_id: str) -> AskRunRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, document_version_id, job_id, question, source_toggles_json,
                   status, answer_type, answer_markdown, citations_json, created_at, completed_at
            FROM ask_runs
            WHERE id = ?
            """,
            (ask_run_id,),
        ).fetchone()
        if row is None:
            return None
        return build_ask_run_record(row)


def create_query_run(payload: QueryRunCreateRequest) -> QueryRunRecord:
    init_repository()
    query_run_id = f"qry-{uuid4().hex[:12]}"
    created_at = now_timestamp()
    normalized_questions = [
        question.strip()
        for question in payload.questions
        if question and question.strip()
    ]
    normalized_document_version_ids = [
        document_version_id.strip()
        for document_version_id in payload.document_version_ids
        if document_version_id and document_version_id.strip()
    ]
    if not normalized_questions:
        raise ValueError("At least one query question is required.")
    if not normalized_document_version_ids:
        raise ValueError("At least one document is required for a query run.")

    with get_connection() as connection:
        job_id = create_job_record(
            connection=connection,
            job_type="query_run",
            status=JobStatus.QUEUED,
            project_id=payload.project_id,
            document_version_id=None,
            export_id=None,
            metadata={"query_run_id": query_run_id, "question_count": len(normalized_questions)},
        )
        connection.execute(
            """
            INSERT INTO query_runs (
                id, project_id, job_id, name, document_version_ids_json, questions_json,
                status, rows_json, created_at, completed_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                query_run_id,
                payload.project_id,
                job_id,
                payload.name.strip() if payload.name else None,
                json.dumps(normalized_document_version_ids),
                json.dumps(normalized_questions),
                JobStatus.QUEUED,
                json.dumps([]),
                created_at,
                None,
            ),
        )
        insert_audit_event(
            connection=connection,
            event_type="query_run.queued",
            actor_surface="web_app",
            project_id=payload.project_id,
            document_version_id=None,
            review_run_id=None,
            suggestion_id=None,
            payload={
                "query_run_id": query_run_id,
                "job_id": job_id,
                "question_count": len(normalized_questions),
                "document_count": len(normalized_document_version_ids),
            },
        )
        row = connection.execute(
            """
            SELECT id, project_id, job_id, name, document_version_ids_json, questions_json,
                   status, rows_json, created_at, completed_at
            FROM query_runs
            WHERE id = ?
            """,
            (query_run_id,),
        ).fetchone()
        return build_query_run_record(row)


def get_query_run(query_run_id: str) -> QueryRunRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, job_id, name, document_version_ids_json, questions_json,
                   status, rows_json, created_at, completed_at
            FROM query_runs
            WHERE id = ?
            """,
            (query_run_id,),
        ).fetchone()
        if row is None:
            return None
        return build_query_run_record(row)


def list_project_query_runs(project_id: str) -> list[QueryRunRecord]:
    init_repository()
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, project_id, job_id, name, document_version_ids_json, questions_json,
                   status, rows_json, created_at, completed_at
            FROM query_runs
            WHERE project_id = ?
            ORDER BY created_at DESC
            """,
            (project_id,),
        ).fetchall()
        return [build_query_run_record(row) for row in rows]


def get_workflow_template(template_id: str) -> WorkflowTemplateRecord | None:
    for template in load_workflow_templates():
        if template.id == template_id:
            return template
    return None


def get_standards_template(template_id: str) -> StandardsTemplateRecord | None:
    for template in load_standards_templates():
        if template.id == template_id:
            return template
    return None


def create_workflow_run(payload: WorkflowRunCreateRequest) -> WorkflowRunRecord:
    init_repository()
    workflow_template = get_workflow_template(payload.workflow_template_id)
    if workflow_template is None:
        raise ValueError("Workflow template not found.")
    variant, _, _ = resolve_workflow_artifact_variant(
        workflow_template,
        payload.artifact_variant_id,
    )

    normalized_document_version_ids = [
        document_version_id.strip()
        for document_version_id in payload.document_version_ids
        if document_version_id and document_version_id.strip()
    ]
    if not normalized_document_version_ids:
        raise ValueError("At least one document is required for a workflow run.")

    workflow_run_id = f"wfr-{uuid4().hex[:12]}"
    created_at = now_timestamp()

    with get_connection() as connection:
        job_id = create_job_record(
            connection=connection,
            job_type="workflow_run",
            status=JobStatus.QUEUED,
            project_id=payload.project_id,
            document_version_id=None,
            export_id=None,
            metadata={
                "workflow_run_id": workflow_run_id,
                "workflow_template_id": payload.workflow_template_id,
                "artifact_variant_id": variant.id if variant else "",
                "question_count": len(workflow_template.questions),
            },
        )
        connection.execute(
            """
            INSERT INTO workflow_runs (
                id, project_id, workflow_template_id, artifact_variant_id, job_id, name, document_version_ids_json,
                status, query_run_id, dd_report_id, rows_json, created_at, completed_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                workflow_run_id,
                payload.project_id,
                payload.workflow_template_id,
                variant.id if variant else None,
                job_id,
                payload.name.strip() if payload.name else workflow_template.name,
                json.dumps(normalized_document_version_ids),
                JobStatus.QUEUED,
                None,
                None,
                json.dumps([]),
                created_at,
                None,
            ),
        )
        insert_audit_event(
            connection=connection,
            event_type="workflow_run.queued",
            actor_surface="web_app",
            project_id=payload.project_id,
            document_version_id=None,
            review_run_id=None,
            suggestion_id=None,
            payload={
                "workflow_run_id": workflow_run_id,
                "workflow_template_id": payload.workflow_template_id,
                "artifact_variant_id": variant.id if variant else None,
                "job_id": job_id,
                "document_count": len(normalized_document_version_ids),
            },
        )
        row = connection.execute(
            """
            SELECT id, project_id, workflow_template_id, artifact_variant_id, job_id, name, document_version_ids_json,
                   status, query_run_id, dd_report_id, rows_json, created_at, completed_at
            FROM workflow_runs
            WHERE id = ?
            """,
            (workflow_run_id,),
        ).fetchone()
        return build_workflow_run_record(row)


def create_standards_run(payload: StandardsRunCreateRequest) -> StandardsRunRecord:
    init_repository()
    standards_template = get_standards_template(payload.standards_template_id)
    if standards_template is None:
        raise ValueError("Standards template not found.")

    selection_text = payload.selection_text.strip()
    if not selection_text and payload.selection_anchor is not None:
        selection_text = payload.selection_anchor.quote.strip()
    if not selection_text:
        raise ValueError("Selection text is required for standards comparison.")

    weak_clauses: list[StandardsWeakClause] = []
    missing_clauses: list[str] = []
    normalized_text = selection_text.lower()
    matched_count = 0

    for clause in standards_template.required_clauses:
        clause_terms = [term.lower() for term in clause.required_terms]
        matched_terms = [term for term in clause_terms if term in normalized_text]
        if matched_terms:
            matched_count += 1
            if len(matched_terms) < len(clause_terms):
                weak_clauses.append(
                    StandardsWeakClause(
                        title=clause.label,
                        severity=clause.severity,
                        explanation=(
                            f"The clause partially matches the house position but is missing "
                            f"{len(clause_terms) - len(matched_terms)} expected concept(s)."
                        ),
                        suggested_fix=clause.recommended_fix,
                    )
                )
        else:
            missing_clauses.append(clause.label)

    total_clause_count = max(len(standards_template.required_clauses), 1)
    weak_penalty = len(weak_clauses) * 0.5
    coverage_score = round(max(((matched_count - weak_penalty) / total_clause_count) * 100, 0), 1)
    standards_run_id = f"sr-{uuid4().hex[:12]}"
    created_at = now_timestamp()

    with get_connection() as connection:
        connection.execute(
            """
            INSERT INTO standards_runs (
                id, project_id, document_version_id, standards_template_id, comparison_mode,
                selection_text, selection_anchor_json, status, coverage_score,
                missing_clauses_json, weak_clauses_json, created_at, completed_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                standards_run_id,
                payload.project_id,
                payload.document_version_id,
                payload.standards_template_id,
                standards_template.comparison_mode,
                selection_text,
                json.dumps(payload.selection_anchor.model_dump()) if payload.selection_anchor else None,
                JobStatus.SUCCEEDED,
                coverage_score,
                json.dumps(missing_clauses),
                json.dumps([clause.model_dump() for clause in weak_clauses]),
                created_at,
                created_at,
            ),
        )
        insert_audit_event(
            connection=connection,
            event_type="standards_run.completed",
            actor_surface="word_addin",
            project_id=payload.project_id,
            document_version_id=payload.document_version_id,
            review_run_id=None,
            suggestion_id=None,
            payload={
                "standards_run_id": standards_run_id,
                "standards_template_id": payload.standards_template_id,
                "coverage_score": coverage_score,
            },
        )
        row = connection.execute(
            """
            SELECT id, project_id, document_version_id, standards_template_id, comparison_mode,
                   status, coverage_score, missing_clauses_json, weak_clauses_json,
                   created_at, completed_at
            FROM standards_runs
            WHERE id = ?
            """,
            (standards_run_id,),
        ).fetchone()
        return build_standards_run_record(row)


def get_workflow_run(workflow_run_id: str) -> WorkflowRunRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, workflow_template_id, artifact_variant_id, job_id, name, document_version_ids_json,
                   status, query_run_id, dd_report_id, rows_json, created_at, completed_at
            FROM workflow_runs
            WHERE id = ?
            """,
            (workflow_run_id,),
        ).fetchone()
        if row is None:
            return None
        return build_workflow_run_record(row)


def list_project_workflow_runs(project_id: str) -> list[WorkflowRunRecord]:
    init_repository()
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, project_id, workflow_template_id, artifact_variant_id, job_id, name, document_version_ids_json,
                   status, query_run_id, dd_report_id, rows_json, created_at, completed_at
            FROM workflow_runs
            WHERE project_id = ?
            ORDER BY created_at DESC
            """,
            (project_id,),
        ).fetchall()
        return [build_workflow_run_record(row) for row in rows]


def get_standards_run(standards_run_id: str) -> StandardsRunRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, document_version_id, standards_template_id, comparison_mode,
                   status, coverage_score, missing_clauses_json, weak_clauses_json,
                   created_at, completed_at
            FROM standards_runs
            WHERE id = ?
            """,
            (standards_run_id,),
        ).fetchone()
        if row is None:
            return None
        return build_standards_run_record(row)


def get_dd_report(dd_report_id: str) -> DdReportRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, workflow_run_id, project_id, query_run_id, memo_markdown,
                   exceptions_list_json, document_summaries_json, created_at
            FROM dd_reports
            WHERE id = ?
            """,
            (dd_report_id,),
        ).fetchone()
        if row is None:
            return None
        return build_dd_report_record(row)


def update_dd_report(dd_report_id: str, payload: DdReportUpdateRequest) -> DdReportRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, workflow_run_id, project_id
            FROM dd_reports
            WHERE id = ?
            """,
            (dd_report_id,),
        ).fetchone()
        if row is None:
            return None

        current_row = connection.execute(
            """
            SELECT memo_markdown, exceptions_list_json, document_summaries_json
            FROM dd_reports
            WHERE id = ?
            """,
            (dd_report_id,),
        ).fetchone()
        insert_dd_report_event(
            connection=connection,
            dd_report_id=dd_report_id,
            action="updated",
            actor_surface="web_app",
            previous_memo_markdown=current_row["memo_markdown"],
            previous_exceptions_list=json.loads(current_row["exceptions_list_json"]),
            previous_document_summaries=[
                DdReportDocumentSummary.model_validate(entry)
                for entry in json.loads(current_row["document_summaries_json"])
            ],
            current_memo_markdown=payload.memo_markdown,
            current_exceptions_list=payload.exceptions_list,
            current_document_summaries=payload.document_summaries,
        )

        connection.execute(
            """
            UPDATE dd_reports
            SET memo_markdown = ?, exceptions_list_json = ?, document_summaries_json = ?
            WHERE id = ?
            """,
            (
                payload.memo_markdown,
                json.dumps(payload.exceptions_list),
                json.dumps([summary.model_dump() for summary in payload.document_summaries]),
                dd_report_id,
            ),
        )
        insert_audit_event(
            connection=connection,
            event_type="dd_report.updated",
            actor_surface="web_app",
            project_id=row["project_id"],
            document_version_id=None,
            review_run_id=None,
            suggestion_id=None,
            payload={"dd_report_id": dd_report_id, "workflow_run_id": row["workflow_run_id"]},
        )
        updated_row = connection.execute(
            """
            SELECT id, workflow_run_id, project_id, query_run_id, memo_markdown,
                   exceptions_list_json, document_summaries_json, created_at
            FROM dd_reports
            WHERE id = ?
            """,
            (dd_report_id,),
        ).fetchone()
        return build_dd_report_record(updated_row)


def export_dd_report_artifact(dd_report_id: str, format: str) -> tuple[bytes | str, str, str] | None:
    report = get_dd_report(dd_report_id)
    if report is None:
        return None
    if format == "memo":
        return report.memo_markdown, "text/markdown", f"{dd_report_id}-memo.md"
    if format == "exceptions":
        content = "\n".join(f"- {item}" for item in report.exceptions_list)
        return content, "text/plain", f"{dd_report_id}-exceptions.txt"
    if format == "docx":
        return export_dd_report_docx(report)
    return None


def update_workflow_run(workflow_run_id: str, payload: WorkflowRunUpdateRequest) -> WorkflowRunRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, workflow_template_id, artifact_variant_id, dd_report_id
            FROM workflow_runs
            WHERE id = ?
            """,
            (workflow_run_id,),
        ).fetchone()
        if row is None:
            return None

        current_rows_row = connection.execute(
            "SELECT rows_json FROM workflow_runs WHERE id = ?",
            (workflow_run_id,),
        ).fetchone()
        insert_workflow_run_event(
            connection=connection,
            workflow_run_id=workflow_run_id,
            action="rows_updated",
            actor_surface="web_app",
            previous_rows=[
                QueryRunRow.model_validate(entry)
                for entry in json.loads(current_rows_row["rows_json"])
            ],
            current_rows=payload.rows,
        )

        connection.execute(
            """
            UPDATE workflow_runs
            SET rows_json = ?
            WHERE id = ?
            """,
            (json.dumps([entry.model_dump() for entry in payload.rows]), workflow_run_id),
        )

        if row["dd_report_id"]:
            workflow_template = get_workflow_template(row["workflow_template_id"])
            if workflow_template is not None:
                current_report_row = connection.execute(
                    """
                    SELECT memo_markdown, exceptions_list_json, document_summaries_json
                    FROM dd_reports
                    WHERE id = ?
                    """,
                    (row["dd_report_id"],),
                ).fetchone()
                rebuilt_report = build_dd_report(
                    workflow_run_id=workflow_run_id,
                    project_id=row["project_id"],
                    template=workflow_template,
                    artifact_variant_id=row["artifact_variant_id"],
                    rows=payload.rows,
                )
                if current_report_row is not None:
                    insert_dd_report_event(
                        connection=connection,
                        dd_report_id=row["dd_report_id"],
                        action="regenerated_from_rows",
                        actor_surface="web_app",
                        previous_memo_markdown=current_report_row["memo_markdown"],
                        previous_exceptions_list=json.loads(
                            current_report_row["exceptions_list_json"]
                        ),
                        previous_document_summaries=[
                            DdReportDocumentSummary.model_validate(entry)
                            for entry in json.loads(
                                current_report_row["document_summaries_json"]
                            )
                        ],
                        current_memo_markdown=rebuilt_report.memo_markdown,
                        current_exceptions_list=rebuilt_report.exceptions_list,
                        current_document_summaries=rebuilt_report.document_summaries,
                    )
                connection.execute(
                    """
                    UPDATE dd_reports
                    SET memo_markdown = ?, exceptions_list_json = ?, document_summaries_json = ?
                    WHERE id = ?
                    """,
                    (
                        rebuilt_report.memo_markdown,
                        json.dumps(rebuilt_report.exceptions_list),
                        json.dumps(
                            [summary.model_dump() for summary in rebuilt_report.document_summaries]
                        ),
                        row["dd_report_id"],
                    ),
                )

        insert_audit_event(
            connection=connection,
            event_type="workflow_run.updated",
            actor_surface="web_app",
            project_id=row["project_id"],
            document_version_id=None,
            review_run_id=None,
            suggestion_id=None,
            payload={"workflow_run_id": workflow_run_id, "row_count": len(payload.rows)},
        )
        updated_row = connection.execute(
            """
            SELECT id, project_id, workflow_template_id, artifact_variant_id, job_id, name, document_version_ids_json,
                   status, query_run_id, dd_report_id, rows_json, created_at, completed_at
            FROM workflow_runs
            WHERE id = ?
            """,
            (workflow_run_id,),
        ).fetchone()
        return build_workflow_run_record(updated_row)


def rerun_workflow_run(
    workflow_run_id: str,
    payload: WorkflowRunRerunRequest,
) -> WorkflowRunRecord | None:
    existing = get_workflow_run(workflow_run_id)
    if existing is None:
        return None
    return create_workflow_run(
        WorkflowRunCreateRequest(
            project_id=existing.project_id,
            workflow_template_id=existing.workflow_template_id,
            document_version_ids=existing.document_version_ids,
            artifact_variant_id=payload.artifact_variant_id or existing.artifact_variant_id,
            name=payload.name or f"Rerun of {existing.name or existing.workflow_template_id}",
        )
    )


def export_query_run_csv(query_run_id: str) -> str | None:
    query_run = get_query_run(query_run_id)
    if query_run is None:
        return None

    buffer = StringIO()
    csv_writer = writer(buffer)
    csv_writer.writerow(["Document", *query_run.questions])
    for row in query_run.rows:
        csv_writer.writerow(
            [
                row.document_name,
                *[cell.answer for cell in row.cells],
            ]
        )
    return buffer.getvalue()


def export_workflow_run_xlsx(workflow_run_id: str) -> tuple[bytes, str, str] | None:
    workflow_run = get_workflow_run(workflow_run_id)
    if workflow_run is None:
        return None
    report = get_dd_report(workflow_run.dd_report_id) if workflow_run.dd_report_id else None
    template = get_workflow_template(workflow_run.workflow_template_id)
    if template is None:
        template = normalize_workflow_template(
            WorkflowTemplateRecord(
                id=workflow_run.workflow_template_id,
                name=workflow_run.name or "Workflow Export",
                version="0.1",
                description="Fallback workflow export template",
                questions=[cell.question for cell in (workflow_run.rows[0].cells if workflow_run.rows else [])],
                exports=["xlsx"],
            )
        )

    workbook = OpenpyxlWorkbook()
    workbook.remove(workbook.active)

    _, workbook_sheets, _ = resolve_workflow_artifact_variant(
        template,
        workflow_run.artifact_variant_id,
    )

    for sheet_template in workbook_sheets:
        title = sheet_template.title[:31]
        if sheet_template.kind == "results":
            build_workflow_results_sheet(workbook, title, workflow_run.rows, sheet_template)
        elif sheet_template.kind == "citations":
            build_workflow_citations_sheet(workbook, title, workflow_run.rows, sheet_template)
        elif sheet_template.kind == "exceptions" and report is not None:
            build_workflow_exceptions_sheet(workbook, title, report)
        elif sheet_template.kind == "document_summaries" and report is not None:
            build_workflow_summaries_sheet(workbook, title, report)
        elif sheet_template.kind == "history":
            build_workflow_history_sheet(workbook, title, workflow_run, report)

    if not workbook.sheetnames:
        build_workflow_results_sheet(
            workbook,
            "Workflow Results",
            workflow_run.rows,
            WorkflowWorkbookSheetTemplate(id="workflow-results", title="Workflow Results", kind="results"),
        )

    output = BytesIO()
    workbook.save(output)
    return (
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        f"{workflow_run_id}-workflow-results.xlsx",
    )


def export_dd_report_docx(report: DdReportRecord) -> tuple[bytes, str, str]:
    workflow_run = get_workflow_run(report.workflow_run_id)
    title = workflow_run.name if workflow_run and workflow_run.name else "Due Diligence Report"
    template = get_workflow_template(workflow_run.workflow_template_id) if workflow_run else None
    document = DocxDocument()
    document.add_heading(title, level=1)

    metadata_table = document.add_table(rows=3, cols=2)
    metadata_table.style = "Table Grid"
    metadata_table.cell(0, 0).text = "Report ID"
    metadata_table.cell(0, 1).text = report.id
    metadata_table.cell(1, 0).text = "Workflow Run"
    metadata_table.cell(1, 1).text = report.workflow_run_id
    metadata_table.cell(2, 0).text = "Generated"
    metadata_table.cell(2, 1).text = report.created_at

    if template is not None:
        _, _, memo_sections = resolve_workflow_artifact_variant(
            template,
            workflow_run.artifact_variant_id if workflow_run else None,
        )
        for section in memo_sections:
            document.add_heading(section.heading, level=2)
            add_docx_section_content(document, section, workflow_run.rows if workflow_run else [], report)
    else:
        for line in report.memo_markdown.splitlines():
            stripped = line.strip()
            if not stripped:
                document.add_paragraph("")
                continue
            if stripped.startswith("# "):
                document.add_heading(stripped[2:], level=1)
            elif stripped.startswith("## "):
                document.add_heading(stripped[3:], level=2)
            elif stripped.startswith("- "):
                document.add_paragraph(stripped[2:], style="List Bullet")
            elif stripped.startswith("_") and stripped.endswith("_"):
                paragraph = document.add_paragraph(stripped.strip("_"))
                if paragraph.runs:
                    paragraph.runs[0].italic = True
            else:
                document.add_paragraph(stripped)

    for section in document.sections:
        section.top_margin = Pt(54)
        section.bottom_margin = Pt(54)

    output = BytesIO()
    document.save(output)
    return (
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        f"{report.id}-memo.docx",
    )


def autosize_openpyxl_sheet(sheet: object, max_width: int = 48) -> None:
    for column in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(
            max(max_length + 2, 14),
            max_width,
        )


def select_row_cells(
    row: QueryRunRow,
    question_terms: list[str],
) -> list[QueryRunCell]:
    if not question_terms:
        return row.cells
    lowered_terms = [term.lower() for term in question_terms]
    return [
        cell for cell in row.cells
        if any(term in cell.question.lower() for term in lowered_terms)
    ]


def build_workflow_results_sheet(
    workbook: OpenpyxlWorkbook,
    title: str,
    rows: list[QueryRunRow],
    sheet_template: WorkflowWorkbookSheetTemplate,
) -> None:
    sheet = workbook.create_sheet(title)
    reference_cells = (
        select_row_cells(rows[0], sheet_template.question_terms) if rows else []
    )
    headers = ["Document", *[cell.question for cell in reference_cells]]
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4F73")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row_index, row in enumerate(rows, start=2):
        sheet.cell(row=row_index, column=1, value=row.document_name)
        filtered_cells = select_row_cells(row, sheet_template.question_terms)
        for column_index, cell_payload in enumerate(filtered_cells, start=2):
            cell = sheet.cell(row=row_index, column=column_index, value=cell_payload.answer)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    autosize_openpyxl_sheet(sheet, max_width=48)
    sheet.freeze_panes = "A2"


def build_workflow_citations_sheet(
    workbook: OpenpyxlWorkbook,
    title: str,
    rows: list[QueryRunRow],
    sheet_template: WorkflowWorkbookSheetTemplate,
) -> None:
    sheet = workbook.create_sheet(title)
    headers = ["Document", "Question", "Answer", "Citation", "Quoted Snippet"]
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6A4C2A")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    row_index = 2
    for row in rows:
        for cell_payload in select_row_cells(row, sheet_template.question_terms):
            citations = cell_payload.citations or [None]
            for citation in citations:
                sheet.cell(row=row_index, column=1, value=row.document_name)
                sheet.cell(row=row_index, column=2, value=cell_payload.question)
                sheet.cell(row=row_index, column=3, value=cell_payload.answer)
                sheet.cell(row=row_index, column=4, value=citation.label if citation else "")
                quote_cell = sheet.cell(
                    row=row_index,
                    column=5,
                    value=citation.quote if citation else "",
                )
                quote_cell.alignment = Alignment(vertical="top", wrap_text=True)
                row_index += 1

    autosize_openpyxl_sheet(sheet, max_width=56)
    sheet.freeze_panes = "A2"


def build_workflow_exceptions_sheet(
    workbook: OpenpyxlWorkbook,
    title: str,
    report: DdReportRecord,
) -> None:
    sheet = workbook.create_sheet(title)
    cell = sheet.cell(row=1, column=1, value="Exceptions")
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="8F3630")
    for row_index, item in enumerate(report.exceptions_list, start=2):
        issue_cell = sheet.cell(row=row_index, column=1, value=item)
        issue_cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize_openpyxl_sheet(sheet, max_width=88)


def build_workflow_summaries_sheet(
    workbook: OpenpyxlWorkbook,
    title: str,
    report: DdReportRecord,
) -> None:
    sheet = workbook.create_sheet(title)
    headers = ["Document", "Summary"]
    colors = "4F6954"
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=colors)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row_index, summary in enumerate(report.document_summaries, start=2):
        sheet.cell(row=row_index, column=1, value=summary.document_name)
        cell = sheet.cell(row=row_index, column=2, value=summary.summary)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    autosize_openpyxl_sheet(sheet, max_width=72)
    sheet.freeze_panes = "A2"


def build_workflow_history_sheet(
    workbook: OpenpyxlWorkbook,
    title: str,
    workflow_run: WorkflowRunRecord,
    report: DdReportRecord | None,
) -> None:
    sheet = workbook.create_sheet(title)
    headers = ["Scope", "Action", "Actor", "Created At", "Summary"]
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3E4C59")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    row_index = 2
    for event in workflow_run.events:
        sheet.cell(row=row_index, column=1, value="Workflow")
        sheet.cell(row=row_index, column=2, value=event.action)
        sheet.cell(row=row_index, column=3, value=event.actor_surface)
        sheet.cell(row=row_index, column=4, value=event.created_at)
        summary_cell = sheet.cell(row=row_index, column=5, value=format_workflow_event_summary(event))
        summary_cell.alignment = Alignment(vertical="top", wrap_text=True)
        row_index += 1
    if report is not None:
        for event in report.events:
            sheet.cell(row=row_index, column=1, value="Report")
            sheet.cell(row=row_index, column=2, value=event.action)
            sheet.cell(row=row_index, column=3, value=event.actor_surface)
            sheet.cell(row=row_index, column=4, value=event.created_at)
            summary_cell = sheet.cell(row=row_index, column=5, value=format_dd_report_event_summary(event))
            summary_cell.alignment = Alignment(vertical="top", wrap_text=True)
            row_index += 1
    autosize_openpyxl_sheet(sheet, max_width=72)
    sheet.freeze_panes = "A2"


def add_docx_section_content(
    document: DocxDocument,
    section: WorkflowMemoSectionTemplate,
    rows: list[QueryRunRow],
    report: DdReportRecord,
) -> None:
    if section.kind == "overview":
        document.add_paragraph(
            f"This workflow reviewed {len(rows)} document(s) and surfaced {len(report.exceptions_list)} exception item(s) requiring reviewer attention."
        )
        return

    if section.kind == "question_rollup":
        added = False
        for row in rows:
            answers = collect_row_answers_by_terms(row, section.question_terms)
            if answers:
                document.add_paragraph(f"{row.document_name}: {' | '.join(answers)}", style="List Bullet")
                added = True
        if not added:
            document.add_paragraph(section.empty_text or "No findings surfaced.", style="List Bullet")
        return

    if section.kind == "document_summaries":
        if report.document_summaries:
            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Document"
            table.rows[0].cells[1].text = "Summary"
            for summary in report.document_summaries:
                row = table.add_row().cells
                row[0].text = summary.document_name
                row[1].text = summary.summary
        else:
            document.add_paragraph(section.empty_text or "No document summaries are available.")
        return

    if section.kind == "exceptions":
        if report.exceptions_list:
            table = document.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "#"
            table.rows[0].cells[1].text = "Issue"
            for index, item in enumerate(report.exceptions_list, start=1):
                row = table.add_row().cells
                row[0].text = str(index)
                row[1].text = item
        else:
            document.add_paragraph(section.empty_text or "No material exceptions were surfaced.")
        return

    if section.kind == "history":
        if report.events:
            for event in report.events:
                paragraph = document.add_paragraph(style="List Bullet")
                run = paragraph.add_run(f"{event.created_at} | {event.action} | {event.actor_surface}: ")
                run.bold = True
                paragraph.add_run(format_dd_report_event_summary(event))
        else:
            document.add_paragraph(section.empty_text or "No report history is available.")
        return

    document.add_paragraph(section.empty_text or "No content is configured for this section.")


def format_workflow_event_summary(event: WorkflowRunEventRecord) -> str:
    summary = event.diff_summary
    document_label = ", ".join(summary.changed_documents[:3]) or "no documents"
    question_label = ", ".join(summary.changed_questions[:3]) or "no questions"
    if len(summary.changed_questions) > 3:
        question_label += f", +{len(summary.changed_questions) - 3} more"
    return (
        f"{summary.changed_cell_count} cell(s) across {summary.changed_row_count} row(s); "
        f"documents: {document_label}; questions: {question_label}"
    )


def format_dd_report_event_summary(event: DdReportEventRecord) -> str:
    summary = event.diff_summary
    parts = [
        "memo changed" if summary.memo_changed else "memo unchanged",
        f"{summary.exception_added_count} exception(s) added",
        f"{summary.exception_removed_count} exception(s) removed",
    ]
    if summary.summary_changed_documents:
        parts.append(
            "summary docs: "
            + ", ".join(summary.summary_changed_documents[:3])
            + (
                f", +{len(summary.summary_changed_documents) - 3} more"
                if len(summary.summary_changed_documents) > 3
                else ""
            )
        )
    return "; ".join(parts)


def create_review_run(payload: ReviewRunCreateRequest) -> ReviewRunRecord:
    init_repository()
    review_run_id = f"rr-{uuid4().hex[:12]}"
    created_at = now_timestamp()
    selection_text = payload.selection_text.strip() or (
        payload.scope.anchor.quote.strip() if payload.scope.anchor else ""
    )

    suggestions, summary = generate_review_suggestions(
        review_run_id=review_run_id,
        project_id=payload.project_id,
        document_version_id=payload.document_version_id,
        selection_text=selection_text,
        markup_settings=payload.markup_settings,
        represented_party=payload.represented_party,
        jurisdiction=payload.jurisdiction,
    )

    with get_connection() as connection:
        connection.execute(
            """
            INSERT INTO review_runs (
                id, project_id, document_version_id, review_type, represented_party,
                jurisdiction, audience, scope_mode, scope_anchor_json, selection_text,
                selection_ooxml, deal_context_json, markup_settings_json, playbook_ids_json,
                status, summary_json, created_at, completed_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                review_run_id,
                payload.project_id,
                payload.document_version_id,
                payload.review_type,
                payload.represented_party,
                payload.jurisdiction,
                payload.audience,
                payload.scope.mode,
                json.dumps(payload.scope.anchor.model_dump() if payload.scope.anchor else None),
                selection_text,
                payload.selection_ooxml,
                json.dumps(payload.deal_context),
                json.dumps(payload.markup_settings.model_dump()),
                json.dumps(payload.playbook_ids),
                "succeeded",
                json.dumps(summary.model_dump()),
                created_at,
                created_at,
            ),
        )
        insert_review_suggestions(
            connection=connection,
            review_run_id=review_run_id,
            suggestions=suggestions,
        )
        insert_audit_event(
            connection=connection,
            event_type="review_run.created",
            actor_surface="word_addin",
            project_id=payload.project_id,
            document_version_id=payload.document_version_id,
            review_run_id=review_run_id,
            suggestion_id=None,
            payload={
                "review_type": str(payload.review_type),
                "scope_mode": payload.scope.mode,
                "suggestion_count": len(suggestions),
            },
        )

        return build_review_run_record(connection, review_run_id)


def get_review_run(review_run_id: str) -> ReviewRunRecord | None:
    init_repository()

    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, document_version_id, review_type, represented_party,
                   jurisdiction, audience, status, summary_json, created_at, completed_at
            FROM review_runs
            WHERE id = ?
            """,
            (review_run_id,),
        ).fetchone()
        if row is None:
            return None

        return build_review_run_record(connection, review_run_id)


def export_review_run_summary(review_run_id: str) -> ReviewRunExportRecord | None:
    init_repository()

    with get_connection() as connection:
        row = connection.execute(
            "SELECT id FROM review_runs WHERE id = ?",
            (review_run_id,),
        ).fetchone()
        if row is None:
            return None

        review_run = build_review_run_record(connection, review_run_id)
        export_record = create_review_run_export_record(connection, review_run)
        create_job_record(
            connection=connection,
            job_type="review_export",
            status=JobStatus.SUCCEEDED,
            project_id=review_run.project_id,
            document_version_id=review_run.document_version_id,
            export_id=export_record.id,
            metadata={
                "review_run_id": review_run.id,
                "format": "markdown",
            },
        )
        insert_audit_event(
            connection=connection,
            event_type="review_run.exported",
            actor_surface="word_addin",
            project_id=review_run.project_id,
            document_version_id=review_run.document_version_id,
            review_run_id=review_run.id,
            suggestion_id=None,
            payload={"export_id": export_record.id, "format": "markdown"},
        )

        return export_record


def create_review_run_export_record(
    connection: object,
    review_run: ReviewRunRecord,
) -> ReviewRunExportRecord:
    exported_at = now_timestamp()
    export_id = f"rre-{uuid4().hex[:12]}"
    summary_markdown = build_review_run_summary_markdown(review_run)

    connection.execute(
        """
        INSERT INTO review_run_exports (
            id, review_run_id, project_id, document_version_id, summary_markdown, exported_at
        )
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            export_id,
            review_run.id,
            review_run.project_id,
            review_run.document_version_id,
            summary_markdown,
            exported_at,
        ),
    )

    return ReviewRunExportRecord(
        id=export_id,
        review_run_id=review_run.id,
        project_id=review_run.project_id,
        document_version_id=review_run.document_version_id,
        summary_markdown=summary_markdown,
        exported_at=exported_at,
    )


def apply_review_suggestion(
    suggestion_id: str,
    payload: ReviewSuggestionApplyRequest,
) -> ReviewRunRecord | None:
    init_repository()
    status = (
        SuggestionStatus.APPLIED_COMMENT
        if payload.mode == "comment"
        else SuggestionStatus.APPLIED_REDLINE
    )
    timestamp = now_timestamp()

    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT review_run_id, supporting_excerpt, citations_json
            FROM review_suggestions
            WHERE id = ?
            """,
            (suggestion_id,),
        ).fetchone()
        if row is None:
            return None

        review_run_id = row["review_run_id"]
        applied_anchor = payload.client_application_result.applied_anchor
        reconciliation_status, reconciliation_note = reconcile_anchor_state(
            mode=payload.mode,
            original_excerpt=resolve_original_anchor_quote(
                supporting_excerpt=row["supporting_excerpt"],
                citations_json=row["citations_json"],
            ),
            applied_anchor=applied_anchor,
        )
        note_parts = []
        if payload.reviewer_note:
            note_parts.append(payload.reviewer_note.strip())
        if payload.client_application_result.message:
            note_parts.append(payload.client_application_result.message.strip())
        reviewer_note = " ".join(part for part in note_parts if part)

        connection.execute(
            """
            UPDATE review_suggestions
            SET status = ?, reviewer_note = ?, applied_at = ?, dismissed_at = NULL,
                latest_anchor_json = ?, anchor_reconciliation_status = ?,
                anchor_reconciliation_note = ?
            WHERE id = ?
            """,
            (
                status,
                reviewer_note or None,
                timestamp,
                json.dumps(applied_anchor.model_dump() if applied_anchor else None),
                reconciliation_status,
                reconciliation_note,
                suggestion_id,
            ),
        )
        insert_review_suggestion_event(
            connection=connection,
            suggestion_id=suggestion_id,
            action=str(status),
            actor_surface="word_addin",
            status_after=status,
            note=payload.reviewer_note,
            client_message=payload.client_application_result.message,
            applied_anchor=applied_anchor,
        )
        review_row = connection.execute(
            "SELECT project_id, document_version_id FROM review_runs WHERE id = ?",
            (review_run_id,),
        ).fetchone()
        if review_row is not None:
            insert_audit_event(
                connection=connection,
                event_type="review_suggestion.applied",
                actor_surface="word_addin",
                project_id=review_row["project_id"],
                document_version_id=review_row["document_version_id"],
                review_run_id=review_run_id,
                suggestion_id=suggestion_id,
                payload={"mode": payload.mode, "status": str(status)},
            )

        return build_review_run_record(connection, review_run_id)


def dismiss_review_suggestion(
    suggestion_id: str,
    payload: ReviewSuggestionDismissRequest,
) -> ReviewRunRecord | None:
    init_repository()
    timestamp = now_timestamp()

    with get_connection() as connection:
        row = connection.execute(
            "SELECT review_run_id FROM review_suggestions WHERE id = ?",
            (suggestion_id,),
        ).fetchone()
        if row is None:
            return None

        review_run_id = row["review_run_id"]
        connection.execute(
            """
            UPDATE review_suggestions
            SET status = ?, reviewer_note = ?, dismissed_at = ?
            WHERE id = ?
            """,
            (
                SuggestionStatus.DISMISSED,
                payload.reason.strip() if payload.reason else None,
                timestamp,
                suggestion_id,
            ),
        )
        insert_review_suggestion_event(
            connection=connection,
            suggestion_id=suggestion_id,
            action="dismissed",
            actor_surface="word_addin",
            status_after=SuggestionStatus.DISMISSED,
            note=payload.reason,
            client_message=None,
            applied_anchor=None,
        )
        review_row = connection.execute(
            "SELECT project_id, document_version_id FROM review_runs WHERE id = ?",
            (review_run_id,),
        ).fetchone()
        if review_row is not None:
            insert_audit_event(
                connection=connection,
                event_type="review_suggestion.dismissed",
                actor_surface="word_addin",
                project_id=review_row["project_id"],
                document_version_id=review_row["document_version_id"],
                review_run_id=review_run_id,
                suggestion_id=suggestion_id,
                payload={"reason": payload.reason.strip() if payload.reason else ""},
            )

        return build_review_run_record(connection, review_run_id)


def mark_review_suggestion_reviewed(
    suggestion_id: str,
    payload: ReviewSuggestionMarkReviewedRequest,
) -> ReviewRunRecord | None:
    init_repository()

    with get_connection() as connection:
        row = connection.execute(
            "SELECT review_run_id FROM review_suggestions WHERE id = ?",
            (suggestion_id,),
        ).fetchone()
        if row is None:
            return None

        review_run_id = row["review_run_id"]
        connection.execute(
            """
            UPDATE review_suggestions
            SET status = ?, reviewer_note = ?
            WHERE id = ?
            """,
            (
                SuggestionStatus.REVIEWED,
                payload.note.strip() if payload.note else None,
                suggestion_id,
            ),
        )
        insert_review_suggestion_event(
            connection=connection,
            suggestion_id=suggestion_id,
            action="reviewed",
            actor_surface="word_addin",
            status_after=SuggestionStatus.REVIEWED,
            note=payload.note,
            client_message=None,
            applied_anchor=None,
        )
        review_row = connection.execute(
            "SELECT project_id, document_version_id FROM review_runs WHERE id = ?",
            (review_run_id,),
        ).fetchone()
        if review_row is not None:
            insert_audit_event(
                connection=connection,
                event_type="review_suggestion.reviewed",
                actor_surface="word_addin",
                project_id=review_row["project_id"],
                document_version_id=review_row["document_version_id"],
                review_run_id=review_run_id,
                suggestion_id=suggestion_id,
                payload={"note": payload.note.strip() if payload.note else ""},
            )

        return build_review_run_record(connection, review_run_id)


def save_review_suggestion_to_playbook(
    suggestion_id: str,
    payload: ReviewSuggestionSaveToPlaybookRequest,
) -> ReviewRunRecord | None:
    init_repository()
    timestamp = now_timestamp()

    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT review_run_id, title, issue_type, severity, reviewer_note,
                   supporting_excerpt, fallback_position_text
            FROM review_suggestions
            WHERE id = ?
            """,
            (suggestion_id,),
        ).fetchone()
        if row is None:
            return None

        review_run_id = row["review_run_id"]
        target_playbook = resolve_target_playbook(payload.playbook_id)
        playbook_id = target_playbook.name if target_playbook is not None else "general-playbook"
        playbook_check_id = resolve_target_playbook_check_id(
            target_playbook=target_playbook,
            requested_check_id=payload.playbook_check_id,
            issue_type=row["issue_type"],
        )
        note = resolve_saved_playbook_note(
            explicit_note=payload.note,
            reviewer_note=row["reviewer_note"],
            supporting_excerpt=row["supporting_excerpt"],
        )
        saved_note_id = f"psn-{uuid4().hex[:12]}"

        connection.execute(
            """
            INSERT INTO playbook_saved_notes (
                id, playbook_id, playbook_check_id, suggestion_id, review_run_id, title, issue_type,
                severity, note, supporting_excerpt, fallback_position_text, created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                saved_note_id,
                playbook_id,
                playbook_check_id,
                suggestion_id,
                review_run_id,
                row["title"],
                row["issue_type"],
                row["severity"],
                note,
                row["supporting_excerpt"],
                row["fallback_position_text"],
                timestamp,
            ),
        )
        connection.execute(
            """
            UPDATE review_suggestions
            SET status = ?, reviewer_note = ?, saved_to_playbook_at = ?,
                saved_playbook_note_id = ?, saved_playbook_id = ?,
                saved_playbook_check_id = ?
            WHERE id = ?
            """,
            (
                SuggestionStatus.SAVED_TO_PLAYBOOK,
                note,
                timestamp,
                saved_note_id,
                playbook_id,
                playbook_check_id,
                suggestion_id,
            ),
        )
        insert_review_suggestion_event(
            connection=connection,
            suggestion_id=suggestion_id,
            action="saved_to_playbook",
            actor_surface="word_addin",
            status_after=SuggestionStatus.SAVED_TO_PLAYBOOK,
            note=note,
            client_message=build_saved_to_playbook_message(playbook_id, playbook_check_id),
            applied_anchor=None,
        )
        review_row = connection.execute(
            "SELECT project_id, document_version_id FROM review_runs WHERE id = ?",
            (review_run_id,),
        ).fetchone()
        if review_row is not None:
            insert_audit_event(
                connection=connection,
                event_type="review_suggestion.saved_to_playbook",
                actor_surface="word_addin",
                project_id=review_row["project_id"],
                document_version_id=review_row["document_version_id"],
                review_run_id=review_run_id,
                suggestion_id=suggestion_id,
                payload={
                    "playbook_id": playbook_id,
                    "playbook_check_id": playbook_check_id or "",
                },
            )

        return build_review_run_record(connection, review_run_id)


def list_saved_playbook_notes() -> list[PlaybookSavedNoteRecord]:
    init_repository()

    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, playbook_id, playbook_check_id, suggestion_id, review_run_id, title, issue_type,
                   severity, note, supporting_excerpt, fallback_position_text, created_at
            FROM playbook_saved_notes
            ORDER BY created_at DESC
            LIMIT 25
            """
        ).fetchall()

        return [
            PlaybookSavedNoteRecord(
                id=row["id"],
                playbook_id=row["playbook_id"],
                playbook_check_id=row["playbook_check_id"],
                suggestion_id=row["suggestion_id"],
                review_run_id=row["review_run_id"],
                title=row["title"],
                issue_type=row["issue_type"],
                severity=SeverityLevel(row["severity"]),
                note=row["note"],
                supporting_excerpt=row["supporting_excerpt"],
                fallback_position_text=row["fallback_position_text"],
                created_at=row["created_at"],
            )
            for row in rows
        ]


def upload_documents(
    workspace_id: str,
    files: list[tuple[str, bytes]],
) -> UploadDocumentsResponse:
    init_repository()

    if workspace_id in SEED_WORKSPACES:
        raise ValueError("Create a non-demo workspace before uploading real files.")

    workspace = get_workspace(workspace_id)
    if workspace is None:
        raise ValueError("Workspace not found.")

    notes: list[str] = []
    playbooks = load_playbooks()
    workspace_dir = UPLOADS_DIR / workspace_id
    workspace_dir.mkdir(parents=True, exist_ok=True)
    project = get_project_by_workspace_id(workspace_id)
    if project is None:
        raise ValueError("Project bridge for workspace not found.")
    written_paths: list[Path] = []

    try:
        with get_connection() as connection:
            for filename, content in files:
                validate_upload(filename)
                file_hash = hashlib.sha256(content).hexdigest()
                duplicate = connection.execute(
                    "SELECT id, name FROM documents WHERE workspace_id = ? AND file_hash = ?",
                    (workspace_id, file_hash),
                ).fetchone()
                if duplicate is not None:
                    notes.append(
                        f"Skipped duplicate upload for {filename}; same contents already stored as {duplicate['name']}."
                    )
                    continue

                document_id = f"doc-{uuid4().hex[:12]}"
                stored_name = f"{document_id}-{sanitize_filename(filename)}"
                file_path = workspace_dir / stored_name
                file_path.write_bytes(content)
                written_paths.append(file_path)

                parsed = parse_document(
                    document_id=document_id,
                    filename=filename,
                    file_path=file_path,
                )
                issues = generate_issues(
                    workspace_id=workspace_id,
                    document=parsed.document,
                    pages=parsed.pages,
                    playbooks=playbooks,
                )

                insert_document(
                    connection=connection,
                    workspace_id=workspace_id,
                    document=parsed.document,
                    extracted_text=parsed.text,
                    file_path=file_path,
                    file_hash=file_hash,
                )
                insert_pages(connection=connection, document_id=document_id, pages=parsed.pages)
                insert_issues(connection=connection, issues=issues)
                document_version_id = f"dv-{uuid4().hex[:12]}"
                insert_document_version(
                    connection=connection,
                    document_version_id=document_version_id,
                    project_id=project.id,
                    source_document_id=document_id,
                    document=parsed.document,
                    extracted_text=parsed.text,
                    file_path=file_path,
                    file_hash=file_hash,
                )
                insert_document_anchors(
                    connection=connection,
                    document_version_id=document_version_id,
                    pages=parsed.pages,
                    clauses=parsed.clauses,
                )
                if parsed.clauses:
                    insert_library_items_from_clauses(
                        connection=connection,
                        project_id=project.id,
                        document_version_id=document_version_id,
                        document_name=parsed.document.name,
                        doc_type=parsed.document.doc_type,
                        governing_law=parsed.document.governing_law,
                        counterparty=parsed.document.counterparty,
                        clauses=parsed.clauses,
                    )
                else:
                    anchor_rows = connection.execute(
                        """
                        SELECT id, document_version_id, page_number, quote, metadata_json
                        FROM document_anchors
                        WHERE document_version_id = ?
                        ORDER BY created_at ASC
                        """,
                        (document_version_id,),
                    ).fetchall()
                    insert_library_items(
                        connection=connection,
                        project_id=project.id,
                        document_version_id=document_version_id,
                        document_name=parsed.document.name,
                        doc_type=parsed.document.doc_type,
                        governing_law=parsed.document.governing_law,
                        counterparty=parsed.document.counterparty,
                        anchor_rows=anchor_rows,
                    )
                create_job_record(
                    connection=connection,
                    job_type="document_ingest",
                    status="succeeded",
                    project_id=project.id,
                    document_version_id=document_version_id,
                    export_id=None,
                    metadata={
                        "filename": filename,
                        "issue_count": len(issues),
                        "source_document_id": document_id,
                    },
                )
                insert_audit_event(
                    connection=connection,
                    event_type="document.uploaded",
                    actor_surface="web_app",
                    project_id=project.id,
                    document_version_id=document_version_id,
                    review_run_id=None,
                    suggestion_id=None,
                    payload={
                        "filename": filename,
                        "issue_count": len(issues),
                        "source_document_id": document_id,
                    },
                )

                notes.append(
                    f"Ingested {filename} and created {len(issues)} first-pass issue(s)."
                )

            connection.execute(
                "UPDATE workspaces SET stage = ?, last_updated = ? WHERE id = ?",
                ("First-pass review ready", now_timestamp(), workspace_id),
            )
            connection.execute(
                "UPDATE projects SET stage = ?, last_updated = ? WHERE id = ?",
                ("First-pass review ready", now_timestamp(), project.id),
            )

            row = connection.execute(
                "SELECT id, name, stage, playbook_names_json, last_updated FROM workspaces WHERE id = ?",
                (workspace_id,),
            ).fetchone()
            updated_workspace = build_workspace_detail(connection, row)
    except Exception:
        for file_path in written_paths:
            file_path.unlink(missing_ok=True)
        raise

    return UploadDocumentsResponse(workspace=updated_workspace, notes=notes)


def upload_documents_to_project(
    project_id: str,
    files: list[tuple[str, bytes]],
) -> CanonicalUploadDocumentsResponse:
    project = get_project(project_id)
    if project is None:
        raise ValueError("Project not found.")

    before_document_ids = {
        record.id for record in list_project_document_versions(project_id)
    }
    before_job_ids = {record.id for record in list_project_jobs(project_id)}
    upload_response = upload_documents(workspace_id=project.workspace_id, files=files)

    updated_project = get_project(project_id)
    if updated_project is None:
        raise ValueError("Project not found after upload.")

    document_versions = [
        record
        for record in list_project_document_versions(project_id)
        if record.id not in before_document_ids
    ]
    jobs = [
        record
        for record in list_project_jobs(project_id)
        if record.id not in before_job_ids
    ]

    return CanonicalUploadDocumentsResponse(
        project=updated_project,
        document_versions=document_versions,
        jobs=jobs,
        notes=upload_response.notes,
    )


def list_issues(
    workspace_id: str,
    severity: SeverityLevel | None = None,
    status: IssueStatus | None = None,
    doc_type: DocumentType | None = None,
    issue_type: str | None = None,
) -> list[IssueRecord]:
    workspace = get_workspace(workspace_id)
    if workspace is None:
        return []

    issues = workspace.issues

    if severity is not None:
        issues = [issue for issue in issues if issue.severity == severity]
    if status is not None:
        issues = [issue for issue in issues if issue.status == status]
    if doc_type is not None:
        document_ids = {
            document.id for document in workspace.documents if document.doc_type == doc_type
        }
        issues = [issue for issue in issues if issue.document_id in document_ids]
    if issue_type:
        issues = [
            issue for issue in issues if issue.issue_type.lower() == issue_type.lower()
        ]

    return issues


def insert_document(
    connection: object,
    workspace_id: str,
    document: DocumentRecord,
    extracted_text: str,
    file_path: Path,
    file_hash: str,
) -> None:
    connection.execute(
        """
        INSERT INTO documents (
            id, workspace_id, name, doc_type, counterparty, effective_date, expiry_date,
            renewal_notice_days, auto_renews, governing_law, file_path, file_hash,
            extracted_text, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            document.id,
            workspace_id,
            document.name,
            document.doc_type,
            document.counterparty,
            document.effective_date,
            document.expiry_date,
            document.renewal_notice_days,
            int(document.auto_renews),
            document.governing_law,
            str(file_path),
            file_hash,
            extracted_text,
            now_timestamp(),
        ),
    )


def insert_document_version(
    connection: object,
    document_version_id: str,
    project_id: str,
    source_document_id: str,
    document: DocumentRecord,
    extracted_text: str,
    file_path: Path,
    file_hash: str,
) -> None:
    connection.execute(
        """
        INSERT INTO document_versions (
            id, project_id, source_document_id, name, source_type, file_path, file_hash,
            extracted_text, status, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            document_version_id,
            project_id,
            source_document_id,
            document.name,
            "upload",
            str(file_path),
            file_hash,
            extracted_text,
            "ingested",
            now_timestamp(),
        ),
    )


def insert_pages(connection: object, document_id: str, pages: list[ParsedPage]) -> None:
    connection.executemany(
        """
        INSERT INTO document_pages (id, document_id, page_number, section_heading, text)
        VALUES (?, ?, ?, ?, ?)
        """,
        [
            (
                f"page-{document_id}-{page.page_number}",
                document_id,
                page.page_number,
                page.section_heading,
                page.text,
            )
            for page in pages
        ],
    )


def insert_document_anchors(
    connection: object,
    document_version_id: str,
    pages: list[ParsedPage],
    clauses: list[ParsedClause] | None = None,
) -> None:
    rows: list[
        tuple[
            str,
            str,
            str,
            int | None,
            str | None,
            int | None,
            int | None,
            str,
            str | None,
            str | None,
            str,
            str,
        ]
    ] = []

    if clauses:
        for clause_index, clause in enumerate(clauses, start=1):
            text = normalize_clause_text(clause.text)
            if not text:
                continue
            rows.append(
                (
                    f"anc-{uuid4().hex[:12]}",
                    document_version_id,
                    "pdf_bbox" if clause.page_number else "chunk",
                    clause.page_number,
                    None,
                    0,
                    len(text),
                    text,
                    hashlib.sha256(text.encode("utf-8")).hexdigest()
                    if text
                    else None,
                    None,
                    json.dumps(
                        {
                            "section_heading": clause.section_heading,
                            "clause_index": clause_index,
                            "source_kind": "parsed_clause",
                        }
                    ),
                    now_timestamp(),
                )
            )
    else:
        for page in pages:
            segments = extract_clause_segments(page.text)
            for segment_index, segment in enumerate(segments, start=1):
                rows.append(
                    (
                        f"anc-{uuid4().hex[:12]}",
                        document_version_id,
                        "pdf_bbox" if page.page_number else "chunk",
                        page.page_number,
                        None,
                        0,
                        len(segment),
                        segment,
                        hashlib.sha256(segment.encode("utf-8")).hexdigest()
                        if segment
                        else None,
                        None,
                        json.dumps(
                            {
                                "section_heading": page.section_heading,
                                "segment_index": segment_index,
                                "source_kind": "anchor_segment",
                            }
                        ),
                        now_timestamp(),
                    )
                )

    connection.executemany(
        """
        INSERT INTO document_anchors (
            id, document_version_id, anchor_type, page_number, paragraph_id,
            char_start, char_end, quote, quote_hash, ooxml_path, metadata_json, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )


def insert_library_items(
    connection: object,
    project_id: str,
    document_version_id: str,
    document_name: str,
    doc_type: DocumentType,
    governing_law: str | None,
    counterparty: str | None,
    anchor_rows: list[object],
) -> None:
    rows = []
    for anchor_row in anchor_rows:
        metadata = json.loads(anchor_row["metadata_json"])
        text = normalize_clause_text(anchor_row["quote"])
        if not text:
            continue
        title = derive_library_item_title(text, metadata.get("section_heading", "Clause"))
        rows.append(
            (
                f"lib-{uuid4().hex[:12]}",
                project_id,
                document_version_id,
                anchor_row["id"],
                document_name,
                title,
                metadata.get("section_heading", "Clause"),
                text,
                doc_type,
                governing_law,
                counterparty,
                anchor_row["page_number"],
                "anchor_segment",
                now_timestamp(),
            )
        )

    connection.executemany(
        """
        INSERT INTO library_items (
            id, project_id, document_version_id, anchor_id, document_name, title,
            section_heading, text, doc_type, governing_law, counterparty, page_number,
            source_kind, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )


def insert_library_items_from_clauses(
    connection: object,
    project_id: str,
    document_version_id: str,
    document_name: str,
    doc_type: DocumentType,
    governing_law: str | None,
    counterparty: str | None,
    clauses: list[ParsedClause],
) -> None:
    anchor_rows = connection.execute(
        """
        SELECT id, document_version_id, page_number, quote, metadata_json
        FROM document_anchors
        WHERE document_version_id = ?
        ORDER BY created_at ASC
        """,
        (document_version_id,),
    ).fetchall()
    anchor_ids_by_quote: dict[str, list[str]] = {}
    for anchor_row in anchor_rows:
        normalized_anchor_quote = normalize_clause_text(anchor_row["quote"])
        if not normalized_anchor_quote:
            continue
        anchor_ids_by_quote.setdefault(normalized_anchor_quote, []).append(anchor_row["id"])
    rows = []
    for clause in clauses:
        text = normalize_clause_text(clause.text)
        if not text:
            continue
        matching_anchor_ids = anchor_ids_by_quote.get(text, [])
        matching_anchor_id = matching_anchor_ids.pop(0) if matching_anchor_ids else None
        if matching_anchor_id is None:
            continue
        rows.append(
            (
                f"lib-{uuid4().hex[:12]}",
                project_id,
                document_version_id,
                matching_anchor_id,
                document_name,
                derive_library_item_title(text, clause.section_heading),
                clause.section_heading,
                text,
                doc_type,
                governing_law,
                counterparty,
                clause.page_number,
                "parsed_clause",
                now_timestamp(),
            )
        )

    connection.executemany(
        """
        INSERT INTO library_items (
            id, project_id, document_version_id, anchor_id, document_name, title,
            section_heading, text, doc_type, governing_law, counterparty, page_number,
            source_kind, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )


def insert_issues(connection: object, issues: list[IssueRecord]) -> None:
    connection.executemany(
        """
        INSERT INTO issues (
            id, workspace_id, document_id, title, issue_type, severity, status,
            summary, reviewer_note, counterparties_json, key_dates_json, citations_json, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                issue.id,
                issue.workspace_id,
                issue.document_id,
                issue.title,
                issue.issue_type,
                issue.severity,
                issue.status,
                issue.summary,
                issue.reviewer_note,
                json.dumps(issue.counterparties),
                json.dumps([entry.model_dump() for entry in issue.key_dates]),
                json.dumps([citation.model_dump() for citation in issue.citations]),
                now_timestamp(),
            )
            for issue in issues
        ],
    )


def insert_review_suggestions(
    connection: object,
    review_run_id: str,
    suggestions: list[ReviewSuggestionRecord],
) -> None:
    connection.executemany(
        """
        INSERT INTO review_suggestions (
            id, review_run_id, anchor_id, title, issue_type, severity, confidence,
            explanation, supporting_excerpt, proposed_comment, proposed_redline_json,
            fallback_position_text, status, reviewer_note, applied_at, dismissed_at,
            saved_to_playbook_at, saved_playbook_note_id, saved_playbook_id,
            saved_playbook_check_id, latest_anchor_json, anchor_reconciliation_status,
            anchor_reconciliation_note, citations_json, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                suggestion.id,
                review_run_id,
                suggestion.anchor_id,
                suggestion.title,
                suggestion.issue_type,
                suggestion.severity,
                suggestion.confidence,
                suggestion.explanation,
                suggestion.supporting_excerpt,
                suggestion.proposed_comment,
                json.dumps(
                    suggestion.proposed_redline.model_dump()
                    if suggestion.proposed_redline is not None
                    else None
                ),
                suggestion.fallback_position_text,
                suggestion.status,
                suggestion.reviewer_note,
                suggestion.applied_at,
                suggestion.dismissed_at,
                suggestion.saved_to_playbook_at,
                suggestion.saved_playbook_note_id,
                suggestion.saved_playbook_id,
                suggestion.saved_playbook_check_id,
                json.dumps(suggestion.latest_anchor.model_dump() if suggestion.latest_anchor else None),
                suggestion.anchor_reconciliation_status,
                suggestion.anchor_reconciliation_note,
                json.dumps([citation.model_dump() for citation in suggestion.citations]),
                now_timestamp(),
            )
            for suggestion in suggestions
        ],
    )


def insert_review_suggestion_event(
    connection: object,
    suggestion_id: str,
    action: str,
    actor_surface: str,
    status_after: SuggestionStatus,
    note: str | None,
    client_message: str | None,
    applied_anchor: DocumentAnchor | None,
) -> None:
    connection.execute(
        """
        INSERT INTO review_suggestion_events (
            id, suggestion_id, action, actor_surface, status_after, note,
            client_message, applied_anchor_json, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            f"rse-{uuid4().hex[:12]}",
            suggestion_id,
            action,
            actor_surface,
            status_after,
            note.strip() if note else None,
            client_message,
            json.dumps(applied_anchor.model_dump() if applied_anchor else None),
            now_timestamp(),
        ),
    )


def create_job_record(
    connection: object,
    job_type: str,
    status: JobStatus | str,
    project_id: str | None,
    document_version_id: str | None,
    export_id: str | None,
    metadata: dict[str, str | int | float | bool | None],
    worker_name: str | None = None,
    error_message: str | None = None,
) -> str:
    job_id = f"job-{uuid4().hex[:12]}"
    timestamp = now_timestamp()
    connection.execute(
        """
        INSERT INTO job_records (
            id, job_type, status, project_id, document_version_id, export_id,
            metadata_json, created_at, started_at, completed_at, worker_name, error_message
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            job_id,
            job_type,
            status,
            project_id,
            document_version_id,
            export_id,
            json.dumps(metadata),
            timestamp,
            timestamp if status != JobStatus.QUEUED else None,
            timestamp if status in {JobStatus.SUCCEEDED, JobStatus.FAILED, JobStatus.CANCELED} else None,
            worker_name,
            error_message,
        ),
    )
    return job_id


def get_job(job_id: str) -> JobRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, job_type, status, project_id, document_version_id, export_id,
                   metadata_json, created_at, started_at, completed_at, worker_name, error_message
            FROM job_records
            WHERE id = ?
            """,
            (job_id,),
        ).fetchone()
        if row is None:
            return None
        return build_job_record(row)


def insert_audit_event(
    connection: object,
    event_type: str,
    actor_surface: str,
    project_id: str | None,
    document_version_id: str | None,
    review_run_id: str | None,
    suggestion_id: str | None,
    payload: dict[str, str | int | float | bool | None],
) -> str:
    audit_id = f"ae-{uuid4().hex[:12]}"
    connection.execute(
        """
        INSERT INTO audit_events (
            id, event_type, actor_surface, project_id, document_version_id,
            review_run_id, suggestion_id, payload_json, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            audit_id,
            event_type,
            actor_surface,
            project_id,
            document_version_id,
            review_run_id,
            suggestion_id,
            json.dumps(payload),
            now_timestamp(),
        ),
    )
    return audit_id


def summarize_workflow_row_changes(
    previous_rows: list[QueryRunRow],
    current_rows: list[QueryRunRow],
) -> WorkflowRunEventSummary:
    previous_by_document = {
        row.document_version_id: row for row in previous_rows
    }
    current_by_document = {
        row.document_version_id: row for row in current_rows
    }

    changed_documents: list[str] = []
    changed_questions: set[str] = set()
    changed_cell_count = 0

    all_document_ids = set(previous_by_document) | set(current_by_document)
    for document_id in sorted(all_document_ids):
        previous_row = previous_by_document.get(document_id)
        current_row = current_by_document.get(document_id)
        row_changed = False

        if previous_row is None and current_row is not None:
            row_changed = True
            changed_cell_count += len(current_row.cells)
            changed_questions.update(cell.question for cell in current_row.cells)
            changed_documents.append(current_row.document_name)
            continue

        if current_row is None and previous_row is not None:
            row_changed = True
            changed_cell_count += len(previous_row.cells)
            changed_questions.update(cell.question for cell in previous_row.cells)
            changed_documents.append(previous_row.document_name)
            continue

        if previous_row is None or current_row is None:
            continue

        previous_cells = {cell.question: cell.answer for cell in previous_row.cells}
        current_cells = {cell.question: cell.answer for cell in current_row.cells}
        all_questions = set(previous_cells) | set(current_cells)
        for question in all_questions:
            if previous_cells.get(question) != current_cells.get(question):
                row_changed = True
                changed_cell_count += 1
                changed_questions.add(question)

        if row_changed:
            changed_documents.append(current_row.document_name)

    return WorkflowRunEventSummary(
        changed_row_count=len(changed_documents),
        changed_cell_count=changed_cell_count,
        changed_documents=changed_documents,
        changed_questions=sorted(changed_questions),
    )


def summarize_dd_report_changes(
    previous_memo_markdown: str,
    current_memo_markdown: str,
    previous_exceptions_list: list[str],
    current_exceptions_list: list[str],
    previous_document_summaries: list[DdReportDocumentSummary],
    current_document_summaries: list[DdReportDocumentSummary],
) -> DdReportEventSummary:
    previous_exceptions = set(previous_exceptions_list)
    current_exceptions = set(current_exceptions_list)
    previous_summaries = {
        summary.document_version_id: summary for summary in previous_document_summaries
    }
    current_summaries = {
        summary.document_version_id: summary for summary in current_document_summaries
    }

    changed_documents: list[str] = []
    for document_id in sorted(set(previous_summaries) | set(current_summaries)):
        previous_summary = previous_summaries.get(document_id)
        current_summary = current_summaries.get(document_id)
        previous_text = previous_summary.summary if previous_summary else None
        current_text = current_summary.summary if current_summary else None
        if previous_text != current_text:
            changed_documents.append(
                (current_summary or previous_summary).document_name
            )

    return DdReportEventSummary(
        memo_changed=previous_memo_markdown.strip() != current_memo_markdown.strip(),
        exception_added_count=len(current_exceptions - previous_exceptions),
        exception_removed_count=len(previous_exceptions - current_exceptions),
        summary_changed_documents=changed_documents,
    )


def insert_workflow_run_event(
    connection: object,
    workflow_run_id: str,
    action: str,
    actor_surface: str,
    previous_rows: list[QueryRunRow],
    current_rows: list[QueryRunRow],
) -> str:
    event_id = f"wfe-{uuid4().hex[:12]}"
    diff_summary = summarize_workflow_row_changes(previous_rows, current_rows)
    connection.execute(
        """
        INSERT INTO workflow_run_events (
            id, workflow_run_id, action, actor_surface, previous_rows_json, diff_summary_json, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            event_id,
            workflow_run_id,
            action,
            actor_surface,
            json.dumps([row.model_dump() for row in previous_rows]),
            json.dumps(diff_summary.model_dump()),
            now_timestamp(),
        ),
    )
    return event_id


def insert_dd_report_event(
    connection: object,
    dd_report_id: str,
    action: str,
    actor_surface: str,
    previous_memo_markdown: str,
    previous_exceptions_list: list[str],
    previous_document_summaries: list[DdReportDocumentSummary],
    current_memo_markdown: str,
    current_exceptions_list: list[str],
    current_document_summaries: list[DdReportDocumentSummary],
) -> str:
    event_id = f"dre-{uuid4().hex[:12]}"
    diff_summary = summarize_dd_report_changes(
        previous_memo_markdown=previous_memo_markdown,
        current_memo_markdown=current_memo_markdown,
        previous_exceptions_list=previous_exceptions_list,
        current_exceptions_list=current_exceptions_list,
        previous_document_summaries=previous_document_summaries,
        current_document_summaries=current_document_summaries,
    )
    connection.execute(
        """
        INSERT INTO dd_report_events (
            id, dd_report_id, action, actor_surface, previous_memo_markdown,
            previous_exceptions_list_json, previous_document_summaries_json, diff_summary_json, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            event_id,
            dd_report_id,
            action,
            actor_surface,
            previous_memo_markdown,
            json.dumps(previous_exceptions_list),
            json.dumps([summary.model_dump() for summary in previous_document_summaries]),
            json.dumps(diff_summary.model_dump()),
            now_timestamp(),
        ),
    )
    return event_id


def list_project_jobs(project_id: str) -> list[JobRecord]:
    init_repository()
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, job_type, status, project_id, document_version_id, export_id,
                   metadata_json, created_at, started_at, completed_at, worker_name, error_message
            FROM job_records
            WHERE project_id = ?
            ORDER BY created_at DESC
            """,
            (project_id,),
        ).fetchall()
        return [build_job_record(row) for row in rows]


def list_audit_events(project_id: str | None = None) -> list[AuditEventRecord]:
    init_repository()
    with get_connection() as connection:
        if project_id:
            rows = connection.execute(
                """
                SELECT id, event_type, actor_surface, project_id, document_version_id,
                       review_run_id, suggestion_id, payload_json, created_at
                FROM audit_events
                WHERE project_id = ?
                ORDER BY created_at DESC
                LIMIT 100
                """,
                (project_id,),
            ).fetchall()
        else:
            rows = connection.execute(
                """
                SELECT id, event_type, actor_surface, project_id, document_version_id,
                       review_run_id, suggestion_id, payload_json, created_at
                FROM audit_events
                ORDER BY created_at DESC
                LIMIT 100
                """
            ).fetchall()
        return [build_audit_event_record(row) for row in rows]


def generate_ask_answer(
    connection: object,
    project_id: str,
    document_version_id: str | None,
    question: str,
    answer_type: AnswerType,
    source_toggles: AskSourceToggles,
    selection_text: str | None,
    selection_anchor: DocumentAnchor | None,
) -> tuple[str, list[CitationRecord]]:
    relevant_citations = collect_ask_citations(
        connection=connection,
        document_version_id=document_version_id,
        question=question,
        selection_text=selection_text,
        selection_anchor=selection_anchor,
    )
    source_excerpt = selection_text.strip() if selection_text else (
        relevant_citations[0].quote if relevant_citations else ""
    )
    normalized = f"{question}\n{source_excerpt}".lower()
    mentions_assignment = "assign" in normalized
    mentions_control = "change of control" in normalized or "control" in normalized
    has_consent = "consent" in normalized or "written consent" in normalized
    has_affiliate_carveout = "affiliate" in normalized or "internal reorganization" in normalized

    if mentions_assignment or mentions_control:
        plain_answer = (
            "The clause appears consent-based and does not include a clear affiliate or change-of-control carve-out."
            if has_consent and not has_affiliate_carveout
            else "The clause appears to allow some transfer flexibility, but the exact carve-outs should still be confirmed against the selected text."
        )
    else:
        plain_answer = (
            "Based on the selected contract text, the answer is anchored in the cited excerpt below."
        )

    public_source_note = ""
    if source_toggles.legal_sources or source_toggles.web_search:
        public_source_note = (
            "\n\nPublic legal-source connectors are not configured in this local slice, so this answer is based only on customer document context."
        )

    if answer_type == AnswerType.CLAUSE:
        answer = (
            "Neither party may assign this Agreement without prior written consent, except to an affiliate or in connection with an internal reorganization, merger, or sale of substantially all assets, provided the assigning party remains responsible for its obligations."
        )
    elif answer_type == AnswerType.CHECKLIST:
        answer = (
            "1. Confirm whether consent is required for direct assignment.\n"
            "2. Check whether indirect transfers or change of control are addressed.\n"
            "3. Add an affiliate or internal reorganization carve-out if the represented party needs transfer flexibility."
        )
    elif answer_type == AnswerType.ISSUE_LIST:
        answer = (
            f"- {plain_answer}\n"
            "- The current wording should be compared against the party's house fallback before markup."
        )
    elif answer_type == AnswerType.COMPARISON_TABLE:
        answer = (
            "| Topic | Current text | Skua note |\n"
            "| --- | --- | --- |\n"
            f"| Assignment | {truncate_text(source_excerpt, 80)} | {plain_answer} |"
        )
    elif answer_type == AnswerType.MEMO:
        answer = (
            "The selected agreement language should be treated as a consent-based assignment restriction. "
            + plain_answer
        )
    else:
        answer = plain_answer

    return f"{answer}{public_source_note}", relevant_citations


def generate_draft_result(
    connection: object,
    project_id: str,
    document_version_id: str | None,
    mode: DraftMode,
    query: str | None,
    instruction: str | None,
    selection_text: str | None,
    selection_anchor: DocumentAnchor | None,
) -> tuple[str, list[CitationRecord], list[DraftLibraryMatch]]:
    normalized_query = (query or instruction or "assignment").strip()
    citations = collect_ask_citations(
        connection=connection,
        document_version_id=document_version_id,
        question=normalized_query,
        selection_text=selection_text,
        selection_anchor=selection_anchor,
    )
    base_clause = (
        selection_text.strip()
        if selection_text and selection_text.strip()
        else "Neither party may assign this Agreement without prior written consent."
    )
    library_results = search_library_items_internal(
        connection=connection,
        project_id=project_id,
        query=normalized_query,
        limit=3,
        document_version_id=document_version_id,
    )
    library_matches = (
        build_draft_library_matches_from_library(
            base_clause=base_clause,
            instruction=instruction,
            items=library_results,
        )
        if library_results
        else build_draft_library_matches(base_clause=base_clause, query=normalized_query)
    )

    if mode == DraftMode.INSTRUCTION and instruction and instruction.strip():
        generated_text = (
            "Drafted from instruction:\n\n"
            + build_instruction_draft(base_clause=base_clause, instruction=instruction.strip())
        )
    elif mode == DraftMode.IMPROVE:
        generated_text = (
            "Improved clause:\n\n"
            + library_matches[0].adjusted_text
        )
    else:
        generated_text = (
            "Library-adjusted clause:\n\n"
            + library_matches[0].adjusted_text
        )

    return generated_text, citations, library_matches


def generate_query_run_rows(
    connection: object,
    project_id: str,
    document_version_ids: list[str],
    questions: list[str],
) -> list[QueryRunRow]:
    rows: list[QueryRunRow] = []
    for document_version_id in document_version_ids:
        document_row = connection.execute(
            """
            SELECT dv.id, dv.name, d.counterparty, d.governing_law, d.effective_date,
                   d.expiry_date, d.renewal_notice_days, d.auto_renews
            FROM document_versions dv
            LEFT JOIN documents d ON d.id = dv.source_document_id
            WHERE dv.id = ? AND dv.project_id = ?
            """,
            (document_version_id, project_id),
        ).fetchone()
        if document_row is None:
            continue

        cells = [
            build_query_run_cell(
                connection=connection,
                document_version_id=document_version_id,
                document_row=document_row,
                question=question,
            )
            for question in questions
        ]
        rows.append(
            QueryRunRow(
                document_version_id=document_version_id,
                document_name=document_row["name"],
                cells=cells,
            )
        )
    return rows


def build_query_run_cell(
    connection: object,
    document_version_id: str,
    document_row: object,
    question: str,
) -> QueryRunCell:
    normalized_question = question.strip()
    citations = collect_ask_citations(
        connection=connection,
        document_version_id=document_version_id,
        question=normalized_question,
        selection_text=None,
        selection_anchor=None,
    )
    lowered = normalized_question.lower()

    if "counterparty" in lowered or "party name" in lowered:
        answer = document_row["counterparty"] or "Not found in extracted metadata."
    elif "governing law" in lowered or "venue" in lowered or "jurisdiction" in lowered:
        answer = document_row["governing_law"] or "Not found in extracted metadata."
    elif "effective date" in lowered:
        answer = document_row["effective_date"] or "Not found in extracted metadata."
    elif "expiry" in lowered or "expiration" in lowered:
        answer = document_row["expiry_date"] or "Not found in extracted metadata."
    elif "auto-renew" in lowered or "auto renew" in lowered or "renew" in lowered:
        if document_row["auto_renews"]:
            notice_days = (
                f" with {document_row['renewal_notice_days']} days' notice"
                if document_row["renewal_notice_days"]
                else ""
            )
            answer = f"Yes{notice_days}."
        else:
            answer = "No auto-renewal language detected."
    elif "assignment" in lowered or "change of control" in lowered:
        answer = summarize_assignment_query(citations)
    elif "termination" in lowered:
        answer = summarize_termination_query(citations)
    else:
        answer = summarize_generic_query(citations)

    return QueryRunCell(question=normalized_question, answer=answer, citations=citations[:2])


def summarize_assignment_query(citations: list[CitationRecord]) -> str:
    if not citations:
        return "No assignment language surfaced from the selected documents."
    quote = citations[0].quote.lower()
    if "without prior written consent" in quote or "consent" in quote:
        if "affiliate" in quote or "change of control" in quote:
            return "Consent appears required, but the clause includes an affiliate or structural-transfer carve-out."
        return "Consent appears required on assignment or change of control."
    if "may assign" in quote or "assign this agreement" in quote:
        return "Assignment language is present, but reviewer confirmation is still recommended."
    return summarize_generic_query(citations)


def summarize_termination_query(citations: list[CitationRecord]) -> str:
    if not citations:
        return "No termination language surfaced from the selected documents."
    return truncate_text(citations[0].quote, 180)


def summarize_generic_query(citations: list[CitationRecord]) -> str:
    if not citations:
        return "No responsive language surfaced from the selected documents."
    return truncate_text(citations[0].quote, 180)


def build_dd_report(
    workflow_run_id: str,
    project_id: str,
    template: WorkflowTemplateRecord,
    artifact_variant_id: str | None,
    rows: list[QueryRunRow],
) -> DdReportRecord:
    created_at = now_timestamp()
    report_id = f"ddr-{uuid4().hex[:12]}"
    _, _, memo_sections = resolve_workflow_artifact_variant(template, artifact_variant_id)
    rendered_template = template.model_copy(update={"memo_sections": memo_sections})
    document_summaries = [
        DdReportDocumentSummary(
            document_version_id=row.document_version_id,
            document_name=row.document_name,
            summary=build_document_summary(row),
        )
        for row in rows
    ]
    exceptions = build_exception_list(rows)
    memo_markdown = build_dd_report_markdown(
        template=rendered_template,
        rows=rows,
        document_summaries=document_summaries,
        exceptions=exceptions,
    )
    return DdReportRecord(
        id=report_id,
        workflow_run_id=workflow_run_id,
        project_id=project_id,
        query_run_id=None,
        memo_markdown=memo_markdown,
        exceptions_list=exceptions,
        document_summaries=document_summaries,
        created_at=created_at,
    )


def build_document_summary(row: QueryRunRow) -> str:
    summary_parts = [cell.answer for cell in row.cells[:3] if cell.answer]
    if not summary_parts:
        return "No material findings surfaced from the current extraction pass."
    return " ".join(summary_parts)


def build_exception_list(rows: list[QueryRunRow]) -> list[str]:
    exceptions: list[str] = []
    for row in rows:
        for cell in row.cells:
            answer = cell.answer.lower()
            question = cell.question.lower()
            if "assignment" in question and "consent appears required" in answer:
                exceptions.append(f"{row.document_name}: assignment or change-of-control consent appears required.")
            if "auto-renew" in question and answer.startswith("yes"):
                exceptions.append(f"{row.document_name}: auto-renewal language appears to be present.")
            if "governing law" in question and cell.answer and "ontario" not in answer:
                exceptions.append(f"{row.document_name}: governing law differs from the Ontario default.")
    deduped: list[str] = []
    seen = set()
    for item in exceptions:
        if item in seen:
            continue
        seen.add(item)
        deduped.append(item)
    return deduped[:12]


def build_dd_report_markdown(
    template: WorkflowTemplateRecord,
    rows: list[QueryRunRow],
    document_summaries: list[DdReportDocumentSummary],
    exceptions: list[str],
    report_events: list[DdReportEventRecord] | None = None,
) -> str:
    lines = [
        f"# {template.name}",
        "",
        f"_Generated at {now_timestamp()}_",
        "",
    ]
    for section in template.memo_sections:
        lines.append(f"## {section.heading}")
        lines.append("")
        lines.extend(
            build_memo_section_lines(
                section=section,
                rows=rows,
                document_summaries=document_summaries,
                exceptions=exceptions,
                report_events=report_events or [],
            )
        )
        lines.append("")
    return "\n".join(lines).strip()


def find_cell_answer(row: QueryRunRow, term: str) -> str:
    for cell in row.cells:
        if term in cell.question.lower():
            return cell.answer
    return "No responsive language surfaced."


def build_memo_section_lines(
    section: WorkflowMemoSectionTemplate,
    rows: list[QueryRunRow],
    document_summaries: list[DdReportDocumentSummary],
    exceptions: list[str],
    report_events: list[DdReportEventRecord],
) -> list[str]:
    if section.kind == "overview":
        return [
            f"This workflow reviewed {len(rows)} document(s) and surfaced {len(exceptions)} exception item(s) requiring reviewer attention."
        ]

    if section.kind == "question_rollup":
        rollup_lines: list[str] = []
        for row in rows:
            answers = collect_row_answers_by_terms(row, section.question_terms)
            if answers:
                rollup_lines.append(f"- {row.document_name}: {' | '.join(answers)}")
        return rollup_lines or [section.empty_text or "- No findings surfaced for this section."]

    if section.kind == "document_summaries":
        if document_summaries:
            return [f"- {summary.document_name}: {summary.summary}" for summary in document_summaries]
        return [section.empty_text or "- No document summaries are available."]

    if section.kind == "exceptions":
        if exceptions:
            return [f"- {item}" for item in exceptions]
        return [section.empty_text or "- No material exceptions were surfaced by the current workflow template."]

    if section.kind == "history":
        if report_events:
            return [
                f"- {event.created_at}: {format_dd_report_event_summary(event)}"
                for event in report_events
            ]
        return [section.empty_text or "- No report history is available yet."]

    return [section.empty_text or "- No content is configured for this section."]


def collect_row_answers_by_terms(row: QueryRunRow, question_terms: list[str]) -> list[str]:
    answers: list[str] = []
    lowered_terms = [term.lower() for term in question_terms]
    for cell in row.cells:
        if any(term in cell.question.lower() for term in lowered_terms):
            answers.append(cell.answer)
    return answers


def build_draft_library_matches(base_clause: str, query: str) -> list[DraftLibraryMatch]:
    normalized_query = query.lower()
    affiliate_adjusted = (
        "Neither party may assign this Agreement without prior written consent, except to an affiliate or in connection with a merger, reorganization, or sale of substantially all assets, provided the assigning party remains responsible for its obligations."
    )
    notice_adjusted = (
        "Neither party may assign this Agreement without prior written consent, except to an affiliate as part of an internal reorganization, on prior written notice to the other party."
    )
    tightened_adjusted = (
        "Neither party may assign this Agreement without prior written consent, and any purported assignment in breach of this section is void, except for an affiliate transfer that does not adversely affect the other party."
    )
    first_preview = affiliate_adjusted if "affiliate" in normalized_query or "assignment" in normalized_query else base_clause
    return [
        DraftLibraryMatch(
            id="lib_assignment_ontario_customer",
            title="Customer MSA assignment clause with affiliate carve-out",
            subtitle="Ontario | 2025 approved form",
            preview=first_preview,
            provenance="Source: 2025 Customer MSA / Maple Software",
            adjusted_text=affiliate_adjusted,
        ),
        DraftLibraryMatch(
            id="lib_assignment_affiliate_only",
            title="Vendor paper fallback - affiliate only carve-out",
            subtitle="Ontario | fallback position",
            preview=notice_adjusted,
            provenance="Source: fallback note / vendor paper playbook",
            adjusted_text=notice_adjusted,
        ),
        DraftLibraryMatch(
            id="lib_assignment_tightened",
            title="Conservative consent-based assignment clause",
            subtitle="Ontario | middle-ground fallback",
            preview=tightened_adjusted,
            provenance="Source: negotiated precedent / 2024 vendor compromise",
            adjusted_text=tightened_adjusted,
        ),
    ]


def build_draft_library_matches_from_library(
    base_clause: str,
    instruction: str | None,
    items: list[LibraryItemRecord],
) -> list[DraftLibraryMatch]:
    matches: list[DraftLibraryMatch] = []
    for item in items[:3]:
        adjusted_text = adjust_library_text(
            base_clause=base_clause,
            source_text=item.text,
            instruction=instruction,
        )
        matches.append(
            DraftLibraryMatch(
                id=item.id,
                title=item.title,
                subtitle=f"{item.doc_type.replace('_', ' ')} | {item.document_name}",
                preview=item.text,
                provenance=(
                    f"Source: {item.document_name}"
                    + (f" / {item.section_heading}" if item.section_heading else "")
                    + (f" / p. {item.page_number}" if item.page_number else "")
                ),
                adjusted_text=adjusted_text,
            )
        )
    return matches


def adjust_library_text(base_clause: str, source_text: str, instruction: str | None) -> str:
    lowered = (instruction or "").lower()
    source = source_text.strip()
    if "affiliate" in lowered or "change of control" in lowered:
        if "affiliate" in source.lower() or "change of control" in source.lower():
            return source
        return build_instruction_draft(base_clause=source or base_clause, instruction=instruction or "")
    if source:
        return source
    return base_clause


def build_instruction_draft(base_clause: str, instruction: str) -> str:
    lowered = instruction.lower()
    if "affiliate" in lowered or "change of control" in lowered:
        return (
            "Neither party may assign this Agreement without prior written consent, except to an affiliate or in connection with a merger, internal reorganization, change of control, or sale of substantially all assets, provided the assigning party remains liable for its obligations."
        )
    if "improve" in lowered or "clarify" in lowered:
        return (
            "Neither party may assign this Agreement without prior written consent, and any assignment in breach of this section is void, except for an affiliate transfer made as part of an internal reorganization on prior written notice to the other party."
        )
    return base_clause


def collect_ask_citations(
    connection: object,
    document_version_id: str | None,
    question: str,
    selection_text: str | None,
    selection_anchor: DocumentAnchor | None,
) -> list[CitationRecord]:
    citations: list[CitationRecord] = []

    if selection_text and selection_text.strip():
        synthetic_anchor = selection_anchor or DocumentAnchor(
            type="word_range",
            quote=selection_text.strip(),
        )
        citations.append(
            CitationRecord(
                document_id=document_version_id or "word-live-document",
                document_version_id=document_version_id or "word-live-document",
                anchor_id=synthetic_anchor.id or "selection-anchor",
                label="Current selection",
                quote=selection_text.strip(),
                page_start=synthetic_anchor.page_number,
                page_end=synthetic_anchor.page_number,
            )
        )

    if not document_version_id:
        return citations

    anchor_rows = connection.execute(
        """
        SELECT id, document_version_id, page_number, quote, metadata_json
        FROM document_anchors
        WHERE document_version_id = ?
        ORDER BY created_at ASC
        """,
        (document_version_id,),
    ).fetchall()
    question_terms = {
        term for term in re.findall(r"[a-z]{4,}", question.lower()) if term not in {"what", "does", "this", "that", "with", "from"}
    }
    scored_rows: list[tuple[int, object]] = []
    for row in anchor_rows:
        quote = row["quote"] or ""
        lower_quote = quote.lower()
        score = sum(1 for term in question_terms if term in lower_quote)
        scored_rows.append((score, row))

    for _, row in sorted(scored_rows, key=lambda item: item[0], reverse=True)[:2]:
        if not row["quote"]:
            continue
        metadata = json.loads(row["metadata_json"])
        citations.append(
            CitationRecord(
                document_id=document_version_id,
                document_version_id=document_version_id,
                anchor_id=row["id"],
                label=metadata.get("section_heading", "Document excerpt"),
                quote=row["quote"],
                page_start=row["page_number"],
                page_end=row["page_number"],
            )
        )

    deduped: list[CitationRecord] = []
    seen = set()
    for citation in citations:
        key = (citation.anchor_id, citation.quote)
        if key in seen:
            continue
        seen.add(key)
        deduped.append(citation)
    return deduped[:3]


def search_library_items_internal(
    connection: object,
    project_id: str,
    query: str,
    limit: int,
    document_version_id: str | None = None,
    doc_type: DocumentType | None = None,
    governing_law: str | None = None,
    counterparty: str | None = None,
    document_name: str | None = None,
) -> list[LibraryItemRecord]:
    if document_version_id:
        rows = connection.execute(
            """
            SELECT id, project_id, document_version_id, anchor_id, document_name, title,
                   section_heading, text, doc_type, governing_law, counterparty,
                   page_number, source_kind, created_at
            FROM library_items
            WHERE project_id = ? AND document_version_id = ?
            ORDER BY created_at DESC
            """,
            (project_id, document_version_id),
        ).fetchall()
    else:
        rows = connection.execute(
            """
            SELECT id, project_id, document_version_id, anchor_id, document_name, title,
                   section_heading, text, doc_type, governing_law, counterparty,
                   page_number, source_kind, created_at
            FROM library_items
            WHERE project_id = ?
            ORDER BY created_at DESC
            """,
            (project_id,),
        ).fetchall()

    normalized_governing_law = governing_law.strip().lower() if governing_law else None
    normalized_counterparty = counterparty.strip().lower() if counterparty else None
    normalized_document_name = document_name.strip().lower() if document_name else None

    terms = {
        term
        for term in re.findall(r"[a-z]{3,}", query.lower())
        if term not in {"the", "and", "for", "with", "from", "this", "that"}
    }
    ranked: list[tuple[float, object]] = []
    for row in rows:
        if doc_type and row["doc_type"] != doc_type:
            continue
        if normalized_governing_law and (
            not row["governing_law"] or normalized_governing_law not in row["governing_law"].lower()
        ):
            continue
        if normalized_counterparty and (
            not row["counterparty"] or normalized_counterparty not in row["counterparty"].lower()
        ):
            continue
        if normalized_document_name and normalized_document_name not in row["document_name"].lower():
            continue
        haystack = f"{row['title']} {row['section_heading']} {row['text']}".lower()
        score = float(sum(1 for term in terms if term in haystack))
        if score == 0 and terms:
            continue
        ranked.append((score, row))

    selected_pairs = sorted(ranked, key=lambda entry: entry[0], reverse=True)[:limit]
    return [
        build_library_item_record(row=row, score=score)
        for score, row in selected_pairs
    ]


def build_workspace_summary(connection: object, row: object) -> WorkspaceSummary:
    workspace_id = row["id"]
    document_count = connection.execute(
        "SELECT COUNT(*) FROM documents WHERE workspace_id = ?",
        (workspace_id,),
    ).fetchone()[0]
    issue_count = connection.execute(
        "SELECT COUNT(*) FROM issues WHERE workspace_id = ?",
        (workspace_id,),
    ).fetchone()[0]
    high_severity_count = connection.execute(
        "SELECT COUNT(*) FROM issues WHERE workspace_id = ? AND severity = ?",
        (workspace_id, SeverityLevel.HIGH),
    ).fetchone()[0]

    return WorkspaceSummary(
        id=workspace_id,
        name=row["name"],
        stage=row["stage"],
        playbook_names=json.loads(row["playbook_names_json"]),
        document_count=document_count,
        issue_count=issue_count,
        high_severity_count=high_severity_count,
        last_updated=row["last_updated"],
    )


def build_project_record(connection: object, row: object) -> ProjectRecord:
    project_id = row["id"]
    document_count = connection.execute(
        "SELECT COUNT(*) FROM document_versions WHERE project_id = ?",
        (project_id,),
    ).fetchone()[0]
    latest_job_row = connection.execute(
        """
        SELECT status
        FROM job_records
        WHERE project_id = ?
        ORDER BY created_at DESC
        LIMIT 1
        """,
        (project_id,),
    ).fetchone()
    return ProjectRecord(
        id=project_id,
        workspace_id=row["workspace_id"],
        name=row["name"],
        stage=row["stage"],
        document_count=document_count,
        latest_job_status=(latest_job_row["status"] if latest_job_row else None),
        created_at=row["created_at"],
        last_updated=row["last_updated"],
    )


def list_project_document_versions(project_id: str) -> list[DocumentVersionRecord]:
    init_repository()
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, project_id, source_document_id, name, source_type, status, file_hash, created_at
            FROM document_versions
            WHERE project_id = ?
            ORDER BY created_at DESC
            """,
            (project_id,),
        ).fetchall()
        return [build_document_version_record(connection, row) for row in rows]


def get_document_version(document_version_id: str) -> DocumentVersionDetail | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, source_document_id, name, source_type, status, file_hash, created_at
            FROM document_versions
            WHERE id = ?
            """,
            (document_version_id,),
        ).fetchone()
        if row is None:
            return None

        anchor_rows = connection.execute(
            """
            SELECT id, document_version_id, anchor_type, page_number, paragraph_id,
                   char_start, char_end, quote, quote_hash, ooxml_path
            FROM document_anchors
            WHERE document_version_id = ?
            ORDER BY created_at ASC
            """,
            (document_version_id,),
        ).fetchall()

        anchors = [
            DocumentAnchor(
                id=anchor_row["id"],
                type=anchor_row["anchor_type"],
                document_version_id=anchor_row["document_version_id"],
                paragraph_id=anchor_row["paragraph_id"],
                char_start=anchor_row["char_start"],
                char_end=anchor_row["char_end"],
                page_number=anchor_row["page_number"],
                quote=anchor_row["quote"],
                quote_hash=anchor_row["quote_hash"],
                ooxml_path=anchor_row["ooxml_path"],
            )
            for anchor_row in anchor_rows
        ]

        return DocumentVersionDetail(
            document_version=build_document_version_record(connection, row),
            anchors=anchors,
        )


def search_library_items(payload: LibrarySearchRequest) -> list[LibraryItemRecord]:
    init_repository()
    limit = payload.limit or 5
    with get_connection() as connection:
        clauses = search_library_items_internal(
            connection=connection,
            project_id=payload.project_id,
            query=payload.query,
            limit=limit,
            document_version_id=payload.document_version_id,
            doc_type=payload.doc_type,
            governing_law=payload.governing_law,
            counterparty=payload.counterparty,
            document_name=payload.document_name,
        )
        return clauses


def ingest_document_version(
    document_version_id: str,
    worker_name: str | None = None,
) -> JobRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, source_document_id
            FROM document_versions
            WHERE id = ?
            """,
            (document_version_id,),
        ).fetchone()
        if row is None:
            return None

        job_id = create_job_record(
            connection=connection,
            job_type="document_reingest",
            status=JobStatus.QUEUED,
            project_id=row["project_id"],
            document_version_id=document_version_id,
            export_id=None,
            metadata={"source_document_id": row["source_document_id"] or ""},
            worker_name=worker_name,
        )
        insert_audit_event(
            connection=connection,
            event_type="document.ingest_requested",
            actor_surface="system",
            project_id=row["project_id"],
            document_version_id=document_version_id,
            review_run_id=None,
            suggestion_id=None,
            payload={"job_id": job_id},
        )
        job_row = connection.execute(
            """
            SELECT id, job_type, status, project_id, document_version_id, export_id,
                   metadata_json, created_at, started_at, completed_at, worker_name, error_message
            FROM job_records
            WHERE id = ?
            """,
            (job_id,),
        ).fetchone()
        return build_job_record(job_row)


def queue_review_run_export(review_run_id: str, worker_name: str | None = None) -> JobRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, project_id, document_version_id
            FROM review_runs
            WHERE id = ?
            """,
            (review_run_id,),
        ).fetchone()
        if row is None:
            return None

        job_id = create_job_record(
            connection=connection,
            job_type="review_export",
            status=JobStatus.QUEUED,
            project_id=row["project_id"],
            document_version_id=row["document_version_id"],
            export_id=None,
            metadata={"review_run_id": review_run_id, "format": "markdown"},
            worker_name=worker_name,
        )
        insert_audit_event(
            connection=connection,
            event_type="review_run.export_queued",
            actor_surface="system",
            project_id=row["project_id"],
            document_version_id=row["document_version_id"],
            review_run_id=review_run_id,
            suggestion_id=None,
            payload={"job_id": job_id},
        )
        job_row = connection.execute(
            """
            SELECT id, job_type, status, project_id, document_version_id, export_id,
                   metadata_json, created_at, started_at, completed_at, worker_name, error_message
            FROM job_records
            WHERE id = ?
            """,
            (job_id,),
        ).fetchone()
        if job_row is None:
            return None
        return build_job_record(job_row)


def claim_next_job(worker_name: str) -> JobRecord | None:
    init_repository()
    with get_connection() as connection:
        connection.execute("BEGIN IMMEDIATE")
        row = connection.execute(
            """
            SELECT id
            FROM job_records
            WHERE status = ?
            ORDER BY created_at ASC
            LIMIT 1
            """,
            (JobStatus.QUEUED,),
        ).fetchone()
        if row is None:
            connection.commit()
            return None

        connection.execute(
            """
            UPDATE job_records
            SET status = ?, started_at = ?, worker_name = ?, error_message = NULL
            WHERE id = ? AND status = ?
            """,
            (JobStatus.RUNNING, now_timestamp(), worker_name, row["id"], JobStatus.QUEUED),
        )
        updated = connection.execute(
            "SELECT changes()"
        ).fetchone()[0]
        if updated == 0:
            connection.commit()
            return None

        job_row = connection.execute(
            """
            SELECT id, job_type, status, project_id, document_version_id, export_id,
                   metadata_json, created_at, started_at, completed_at, worker_name, error_message
            FROM job_records
            WHERE id = ?
            """,
            (row["id"],),
        ).fetchone()
        connection.commit()
        return build_job_record(job_row)


def complete_job(
    job_id: str,
    status: JobStatus,
    metadata_updates: dict[str, str | int | float | bool | None] | None = None,
    export_id: str | None = None,
    error_message: str | None = None,
) -> JobRecord | None:
    init_repository()
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, metadata_json
            FROM job_records
            WHERE id = ?
            """,
            (job_id,),
        ).fetchone()
        if row is None:
            return None

        metadata = json.loads(row["metadata_json"])
        if metadata_updates:
            metadata.update(metadata_updates)

        connection.execute(
            """
            UPDATE job_records
            SET status = ?, metadata_json = ?, export_id = COALESCE(?, export_id),
                completed_at = ?, error_message = ?
            WHERE id = ?
            """,
            (
                status,
                json.dumps(metadata),
                export_id,
                now_timestamp(),
                error_message,
                job_id,
            ),
        )
        updated_row = connection.execute(
            """
            SELECT id, job_type, status, project_id, document_version_id, export_id,
                   metadata_json, created_at, started_at, completed_at, worker_name, error_message
            FROM job_records
            WHERE id = ?
            """,
            (job_id,),
        ).fetchone()
        if updated_row is None:
            return None
        return build_job_record(updated_row)


def process_job(job_id: str, worker_name: str) -> JobRecord | None:
    init_repository()
    job = get_job(job_id)
    if job is None:
        return None

    if job.job_type == "workflow_run":
        return process_workflow_run_job(job, worker_name)
    if job.job_type == "query_run":
        return process_query_run_job(job, worker_name)
    if job.job_type == "draft_run":
        return process_draft_run_job(job, worker_name)
    if job.job_type == "ask_run":
        return process_ask_run_job(job, worker_name)
    if job.job_type == "document_reingest":
        return process_document_reingest_job(job, worker_name)
    if job.job_type == "review_export":
        return process_review_export_job(job, worker_name)

    return complete_job(
        job_id=job.id,
        status=JobStatus.FAILED,
        error_message=f"Unsupported job type: {job.job_type}",
    )


def process_next_job(worker_name: str) -> JobRecord | None:
    job = claim_next_job(worker_name)
    if job is None:
        return None

    return process_job(job.id, worker_name)


def process_document_reingest_job(job: JobRecord, worker_name: str) -> JobRecord | None:
    try:
        with get_connection() as connection:
            row = connection.execute(
                """
                SELECT id, project_id, source_document_id
                FROM document_versions
                WHERE id = ?
                """,
                (job.document_version_id,),
            ).fetchone()
            if row is None:
                return complete_job(
                    job.id,
                    status=JobStatus.FAILED,
                    error_message="Document version not found",
                )

            connection.execute(
                "UPDATE document_versions SET status = ? WHERE id = ?",
                ("ingested", row["id"]),
            )
            insert_audit_event(
                connection=connection,
                event_type="document.ingested",
                actor_surface=worker_name,
                project_id=row["project_id"],
                document_version_id=row["id"],
                review_run_id=None,
                suggestion_id=None,
                payload={"job_id": job.id},
            )

        return complete_job(
            job.id,
            status=JobStatus.SUCCEEDED,
            metadata_updates={"result": "document_version_marked_ingested"},
        )
    except Exception as error:  # pragma: no cover - defensive failure path
        return complete_job(
            job.id,
            status=JobStatus.FAILED,
            error_message=str(error),
        )


def process_review_export_job(job: JobRecord, worker_name: str) -> JobRecord | None:
    try:
        review_run_id = str(job.metadata.get("review_run_id") or "")
        if not review_run_id:
            return complete_job(
                job.id,
                status=JobStatus.FAILED,
                error_message="Missing review_run_id metadata",
            )

        with get_connection() as connection:
            review_row = connection.execute(
                "SELECT id FROM review_runs WHERE id = ?",
                (review_run_id,),
            ).fetchone()
            if review_row is None:
                return complete_job(
                    job.id,
                    status=JobStatus.FAILED,
                    error_message="Review run not found",
                )

            review_run = build_review_run_record(connection, review_run_id)
            export_record = create_review_run_export_record(connection, review_run)
            insert_audit_event(
                connection=connection,
                event_type="review_run.exported",
                actor_surface=worker_name,
                project_id=review_run.project_id,
                document_version_id=review_run.document_version_id,
                review_run_id=review_run.id,
                suggestion_id=None,
                payload={"export_id": export_record.id, "job_id": job.id, "format": "markdown"},
            )

        return complete_job(
            job.id,
            status=JobStatus.SUCCEEDED,
            export_id=export_record.id,
            metadata_updates={"export_id": export_record.id, "format": "markdown"},
        )
    except Exception as error:  # pragma: no cover - defensive failure path
        return complete_job(
            job.id,
            status=JobStatus.FAILED,
            error_message=str(error),
        )


def process_query_run_job(job: JobRecord, worker_name: str) -> JobRecord | None:
    try:
        query_run_id = str(job.metadata.get("query_run_id") or "")
        if not query_run_id:
            return complete_job(
                job.id,
                status=JobStatus.FAILED,
                error_message="Missing query_run_id metadata",
            )

        with get_connection() as connection:
            query_row = connection.execute(
                """
                SELECT id, project_id, job_id, name, document_version_ids_json, questions_json
                FROM query_runs
                WHERE id = ?
                """,
                (query_run_id,),
            ).fetchone()
            if query_row is None:
                return complete_job(
                    job.id,
                    status=JobStatus.FAILED,
                    error_message="Query run not found",
                )

            rows = generate_query_run_rows(
                connection=connection,
                project_id=query_row["project_id"],
                document_version_ids=json.loads(query_row["document_version_ids_json"]),
                questions=json.loads(query_row["questions_json"]),
            )
            completed_at = now_timestamp()
            connection.execute(
                """
                UPDATE query_runs
                SET status = ?, rows_json = ?, completed_at = ?
                WHERE id = ?
                """,
                (
                    JobStatus.SUCCEEDED,
                    json.dumps([row.model_dump() for row in rows]),
                    completed_at,
                    query_run_id,
                ),
            )
            insert_audit_event(
                connection=connection,
                event_type="query_run.completed",
                actor_surface=worker_name,
                project_id=query_row["project_id"],
                document_version_id=None,
                review_run_id=None,
                suggestion_id=None,
                payload={"query_run_id": query_run_id, "job_id": job.id, "row_count": len(rows)},
            )

        return complete_job(
            job.id,
            status=JobStatus.SUCCEEDED,
            metadata_updates={"query_run_id": query_run_id},
        )
    except Exception as error:  # pragma: no cover - defensive failure path
        return complete_job(
            job.id,
            status=JobStatus.FAILED,
            error_message=str(error),
        )


def process_workflow_run_job(job: JobRecord, worker_name: str) -> JobRecord | None:
    try:
        workflow_run_id = str(job.metadata.get("workflow_run_id") or "")
        workflow_template_id = str(job.metadata.get("workflow_template_id") or "")
        if not workflow_run_id:
            return complete_job(
                job.id,
                status=JobStatus.FAILED,
                error_message="Missing workflow_run_id metadata",
            )

        workflow_template = get_workflow_template(workflow_template_id)
        if workflow_template is None:
            return complete_job(
                job.id,
                status=JobStatus.FAILED,
                error_message="Workflow template not found",
            )

        with get_connection() as connection:
            workflow_row = connection.execute(
                """
                SELECT id, project_id, workflow_template_id, artifact_variant_id, job_id, name, document_version_ids_json
                FROM workflow_runs
                WHERE id = ?
                """,
                (workflow_run_id,),
            ).fetchone()
            if workflow_row is None:
                return complete_job(
                    job.id,
                    status=JobStatus.FAILED,
                    error_message="Workflow run not found",
                )

            document_version_ids = json.loads(workflow_row["document_version_ids_json"])
            rows = generate_query_run_rows(
                connection=connection,
                project_id=workflow_row["project_id"],
                document_version_ids=document_version_ids,
                questions=workflow_template.questions,
            )
            dd_report = build_dd_report(
                workflow_run_id=workflow_run_id,
                project_id=workflow_row["project_id"],
                template=workflow_template,
                artifact_variant_id=workflow_row["artifact_variant_id"],
                rows=rows,
            )
            completed_at = now_timestamp()
            connection.execute(
                """
                INSERT INTO dd_reports (
                    id, workflow_run_id, project_id, query_run_id, memo_markdown,
                    exceptions_list_json, document_summaries_json, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    dd_report.id,
                    dd_report.workflow_run_id,
                    dd_report.project_id,
                    dd_report.query_run_id,
                    dd_report.memo_markdown,
                    json.dumps(dd_report.exceptions_list),
                    json.dumps([summary.model_dump() for summary in dd_report.document_summaries]),
                    dd_report.created_at,
                ),
            )
            connection.execute(
                """
                UPDATE workflow_runs
                SET status = ?, dd_report_id = ?, rows_json = ?, completed_at = ?
                WHERE id = ?
                """,
                (
                    JobStatus.SUCCEEDED,
                    dd_report.id,
                    json.dumps([row.model_dump() for row in rows]),
                    completed_at,
                    workflow_run_id,
                ),
            )
            insert_audit_event(
                connection=connection,
                event_type="workflow_run.completed",
                actor_surface=worker_name,
                project_id=workflow_row["project_id"],
                document_version_id=None,
                review_run_id=None,
                suggestion_id=None,
                payload={
                    "workflow_run_id": workflow_run_id,
                    "workflow_template_id": workflow_template.id,
                    "dd_report_id": dd_report.id,
                    "job_id": job.id,
                },
            )

        return complete_job(
            job.id,
            status=JobStatus.SUCCEEDED,
            metadata_updates={"workflow_run_id": workflow_run_id, "dd_report_id": dd_report.id},
        )
    except Exception as error:  # pragma: no cover - defensive failure path
        return complete_job(
            job.id,
            status=JobStatus.FAILED,
            error_message=str(error),
        )


def process_ask_run_job(job: JobRecord, worker_name: str) -> JobRecord | None:
    try:
        ask_run_id = str(job.metadata.get("ask_run_id") or "")
        if not ask_run_id:
            return complete_job(
                job.id,
                status=JobStatus.FAILED,
                error_message="Missing ask_run_id metadata",
            )

        with get_connection() as connection:
            ask_row = connection.execute(
                """
                SELECT id, project_id, document_version_id, job_id, selection_anchor_id,
                       selection_text, selection_anchor_json, question, answer_type,
                       source_toggles_json
                FROM ask_runs
                WHERE id = ?
                """,
                (ask_run_id,),
            ).fetchone()
            if ask_row is None:
                return complete_job(
                    job.id,
                    status=JobStatus.FAILED,
                    error_message="Ask run not found",
                )

            answer_markdown, citations = generate_ask_answer(
                connection=connection,
                project_id=ask_row["project_id"],
                document_version_id=ask_row["document_version_id"],
                question=ask_row["question"],
                answer_type=AnswerType(ask_row["answer_type"]),
                source_toggles=AskSourceToggles.model_validate(
                    json.loads(ask_row["source_toggles_json"])
                ),
                selection_text=ask_row["selection_text"],
                selection_anchor=(
                    DocumentAnchor.model_validate(json.loads(ask_row["selection_anchor_json"]))
                    if ask_row["selection_anchor_json"]
                    and json.loads(ask_row["selection_anchor_json"]) is not None
                    else None
                ),
            )
            completed_at = now_timestamp()
            connection.execute(
                """
                UPDATE ask_runs
                SET status = ?, answer_markdown = ?, citations_json = ?, completed_at = ?
                WHERE id = ?
                """,
                (
                    JobStatus.SUCCEEDED,
                    answer_markdown,
                    json.dumps([citation.model_dump() for citation in citations]),
                    completed_at,
                    ask_run_id,
                ),
            )
            insert_audit_event(
                connection=connection,
                event_type="ask_run.completed",
                actor_surface=worker_name,
                project_id=ask_row["project_id"],
                document_version_id=ask_row["document_version_id"],
                review_run_id=None,
                suggestion_id=None,
                payload={"ask_run_id": ask_run_id, "job_id": job.id},
            )

        return complete_job(
            job.id,
            status=JobStatus.SUCCEEDED,
            metadata_updates={"ask_run_id": ask_run_id},
        )
    except Exception as error:  # pragma: no cover - defensive failure path
        return complete_job(
            job.id,
            status=JobStatus.FAILED,
            error_message=str(error),
        )


def process_draft_run_job(job: JobRecord, worker_name: str) -> JobRecord | None:
    try:
        draft_run_id = str(job.metadata.get("draft_run_id") or "")
        if not draft_run_id:
            return complete_job(
                job.id,
                status=JobStatus.FAILED,
                error_message="Missing draft_run_id metadata",
            )

        with get_connection() as connection:
            draft_row = connection.execute(
                """
                SELECT id, project_id, document_version_id, job_id, mode, query, instruction,
                       selection_text, selection_anchor_json
                FROM draft_runs
                WHERE id = ?
                """,
                (draft_run_id,),
            ).fetchone()
            if draft_row is None:
                return complete_job(
                    job.id,
                    status=JobStatus.FAILED,
                    error_message="Draft run not found",
                )

            generated_text, citations, library_matches = generate_draft_result(
                connection=connection,
                project_id=draft_row["project_id"],
                document_version_id=draft_row["document_version_id"],
                mode=DraftMode(draft_row["mode"]),
                query=draft_row["query"],
                instruction=draft_row["instruction"],
                selection_text=draft_row["selection_text"],
                selection_anchor=(
                    DocumentAnchor.model_validate(json.loads(draft_row["selection_anchor_json"]))
                    if draft_row["selection_anchor_json"]
                    and json.loads(draft_row["selection_anchor_json"]) is not None
                    else None
                ),
            )
            completed_at = now_timestamp()
            connection.execute(
                """
                UPDATE draft_runs
                SET status = ?, generated_text = ?, citations_json = ?, library_matches_json = ?, completed_at = ?
                WHERE id = ?
                """,
                (
                    JobStatus.SUCCEEDED,
                    generated_text,
                    json.dumps([citation.model_dump() for citation in citations]),
                    json.dumps([match.model_dump() for match in library_matches]),
                    completed_at,
                    draft_run_id,
                ),
            )
            insert_audit_event(
                connection=connection,
                event_type="draft_run.completed",
                actor_surface=worker_name,
                project_id=draft_row["project_id"],
                document_version_id=draft_row["document_version_id"],
                review_run_id=None,
                suggestion_id=None,
                payload={"draft_run_id": draft_run_id, "job_id": job.id},
            )

        return complete_job(
            job.id,
            status=JobStatus.SUCCEEDED,
            metadata_updates={"draft_run_id": draft_run_id},
        )
    except Exception as error:  # pragma: no cover - defensive failure path
        return complete_job(
            job.id,
            status=JobStatus.FAILED,
            error_message=str(error),
        )


def build_document_version_record(connection: object, row: object) -> DocumentVersionRecord:
    anchor_count = connection.execute(
        "SELECT COUNT(*) FROM document_anchors WHERE document_version_id = ?",
        (row["id"],),
    ).fetchone()[0]
    return DocumentVersionRecord(
        id=row["id"],
        project_id=row["project_id"],
        source_document_id=row["source_document_id"],
        name=row["name"],
        source_type=row["source_type"],
        status=row["status"],
        file_hash=row["file_hash"],
        anchor_count=anchor_count,
        created_at=row["created_at"],
    )


def build_ask_run_record(row: object) -> AskRunRecord:
    return AskRunRecord(
        id=row["id"],
        project_id=row["project_id"],
        document_version_id=row["document_version_id"],
        job_id=row["job_id"],
        question=row["question"],
        source_toggles=AskSourceToggles.model_validate(json.loads(row["source_toggles_json"])),
        status=JobStatus(row["status"]),
        answer_type=AnswerType(row["answer_type"]),
        answer_markdown=row["answer_markdown"],
        citations=[
            CitationRecord.model_validate(entry)
            for entry in json.loads(row["citations_json"])
        ],
        created_at=row["created_at"],
        completed_at=row["completed_at"],
    )


def build_draft_run_record(row: object) -> DraftRunRecord:
    return DraftRunRecord(
        id=row["id"],
        project_id=row["project_id"],
        document_version_id=row["document_version_id"],
        job_id=row["job_id"],
        mode=DraftMode(row["mode"]),
        query=row["query"],
        instruction=row["instruction"],
        status=JobStatus(row["status"]),
        generated_text=row["generated_text"],
        citations=[
            CitationRecord.model_validate(entry)
            for entry in json.loads(row["citations_json"])
        ],
        library_matches=[
            DraftLibraryMatch.model_validate(entry)
            for entry in json.loads(row["library_matches_json"])
        ],
        created_at=row["created_at"],
        completed_at=row["completed_at"],
    )


def build_query_run_record(row: object) -> QueryRunRecord:
    return QueryRunRecord(
        id=row["id"],
        project_id=row["project_id"],
        job_id=row["job_id"],
        name=row["name"],
        document_version_ids=json.loads(row["document_version_ids_json"]),
        questions=json.loads(row["questions_json"]),
        status=JobStatus(row["status"]),
        rows=[
            QueryRunRow.model_validate(entry)
            for entry in json.loads(row["rows_json"])
        ],
        created_at=row["created_at"],
        completed_at=row["completed_at"],
    )


def build_workflow_run_record(row: object) -> WorkflowRunRecord:
    with get_connection() as connection:
        event_rows = connection.execute(
            """
            SELECT id, workflow_run_id, action, actor_surface, previous_rows_json,
                   diff_summary_json, created_at
            FROM workflow_run_events
            WHERE workflow_run_id = ?
            ORDER BY created_at DESC
            """,
            (row["id"],),
        ).fetchall()
    return WorkflowRunRecord(
        id=row["id"],
        project_id=row["project_id"],
        workflow_template_id=row["workflow_template_id"],
        artifact_variant_id=row["artifact_variant_id"],
        job_id=row["job_id"],
        name=row["name"],
        document_version_ids=json.loads(row["document_version_ids_json"]),
        status=JobStatus(row["status"]),
        query_run_id=row["query_run_id"],
        dd_report_id=row["dd_report_id"],
        rows=[QueryRunRow.model_validate(entry) for entry in json.loads(row["rows_json"])],
        events=[
            WorkflowRunEventRecord(
                id=event_row["id"],
                workflow_run_id=event_row["workflow_run_id"],
                action=event_row["action"],
                actor_surface=event_row["actor_surface"],
                previous_rows=[
                    QueryRunRow.model_validate(entry)
                    for entry in json.loads(event_row["previous_rows_json"])
                ],
                diff_summary=WorkflowRunEventSummary.model_validate(
                    json.loads(event_row["diff_summary_json"] or "{}")
                    or {
                        "changed_row_count": 0,
                        "changed_cell_count": 0,
                        "changed_documents": [],
                        "changed_questions": [],
                    }
                ),
                created_at=event_row["created_at"],
            )
            for event_row in event_rows
        ],
        created_at=row["created_at"],
        completed_at=row["completed_at"],
    )


def build_standards_run_record(row: object) -> StandardsRunRecord:
    return StandardsRunRecord(
        id=row["id"],
        project_id=row["project_id"],
        document_version_id=row["document_version_id"],
        standards_template_id=row["standards_template_id"],
        comparison_mode=row["comparison_mode"],
        status=JobStatus(row["status"]),
        coverage_score=row["coverage_score"],
        missing_clauses=json.loads(row["missing_clauses_json"]),
        weak_clauses=[
            StandardsWeakClause.model_validate(entry)
            for entry in json.loads(row["weak_clauses_json"])
        ],
        created_at=row["created_at"],
        completed_at=row["completed_at"],
    )


def build_dd_report_record(row: object) -> DdReportRecord:
    with get_connection() as connection:
        event_rows = connection.execute(
            """
            SELECT id, dd_report_id, action, actor_surface, previous_memo_markdown,
                   previous_exceptions_list_json, previous_document_summaries_json,
                   diff_summary_json, created_at
            FROM dd_report_events
            WHERE dd_report_id = ?
            ORDER BY created_at DESC
            """,
            (row["id"],),
        ).fetchall()
    return DdReportRecord(
        id=row["id"],
        workflow_run_id=row["workflow_run_id"],
        project_id=row["project_id"],
        query_run_id=row["query_run_id"],
        memo_markdown=row["memo_markdown"],
        exceptions_list=json.loads(row["exceptions_list_json"]),
        document_summaries=[
            DdReportDocumentSummary.model_validate(entry)
            for entry in json.loads(row["document_summaries_json"])
        ],
        events=[
            DdReportEventRecord(
                id=event_row["id"],
                dd_report_id=event_row["dd_report_id"],
                action=event_row["action"],
                actor_surface=event_row["actor_surface"],
                previous_memo_markdown=event_row["previous_memo_markdown"],
                previous_exceptions_list=json.loads(event_row["previous_exceptions_list_json"]),
                previous_document_summaries=[
                    DdReportDocumentSummary.model_validate(entry)
                    for entry in json.loads(event_row["previous_document_summaries_json"])
                ],
                diff_summary=DdReportEventSummary.model_validate(
                    json.loads(event_row["diff_summary_json"] or "{}")
                    or {
                        "memo_changed": False,
                        "exception_added_count": 0,
                        "exception_removed_count": 0,
                        "summary_changed_documents": [],
                    }
                ),
                created_at=event_row["created_at"],
            )
            for event_row in event_rows
        ],
        created_at=row["created_at"],
    )


def build_library_item_record(row: object, score: float | None = None) -> LibraryItemRecord:
    return LibraryItemRecord(
        id=row["id"],
        project_id=row["project_id"],
        document_version_id=row["document_version_id"],
        anchor_id=row["anchor_id"],
        document_name=row["document_name"],
        title=row["title"],
        section_heading=row["section_heading"],
        text=row["text"],
        doc_type=DocumentType(row["doc_type"]),
        governing_law=row["governing_law"],
        counterparty=row["counterparty"],
        page_number=row["page_number"],
        source_kind=row["source_kind"],
        score=score,
        created_at=row["created_at"],
    )


def build_job_record(row: object) -> JobRecord:
    return JobRecord(
        id=row["id"],
        job_type=row["job_type"],
        status=JobStatus(row["status"]),
        project_id=row["project_id"],
        document_version_id=row["document_version_id"],
        export_id=row["export_id"],
        metadata=json.loads(row["metadata_json"]),
        created_at=row["created_at"],
        started_at=row["started_at"],
        completed_at=row["completed_at"],
        worker_name=row["worker_name"],
        error_message=row["error_message"],
    )


def build_audit_event_record(row: object) -> AuditEventRecord:
    return AuditEventRecord(
        id=row["id"],
        event_type=row["event_type"],
        actor_surface=row["actor_surface"],
        project_id=row["project_id"],
        document_version_id=row["document_version_id"],
        review_run_id=row["review_run_id"],
        suggestion_id=row["suggestion_id"],
        payload=json.loads(row["payload_json"]),
        created_at=row["created_at"],
    )


def build_review_run_record(connection: object, review_run_id: str) -> ReviewRunRecord:
    row = connection.execute(
        """
        SELECT id, project_id, document_version_id, review_type, represented_party,
               jurisdiction, audience, status, summary_json, created_at, completed_at
        FROM review_runs
        WHERE id = ?
        """,
        (review_run_id,),
    ).fetchone()
    suggestion_rows = connection.execute(
        """
        SELECT id, review_run_id, anchor_id, title, issue_type, severity, confidence,
               explanation, supporting_excerpt, proposed_comment, proposed_redline_json,
               fallback_position_text, status, reviewer_note, applied_at, dismissed_at,
               saved_to_playbook_at, saved_playbook_note_id, saved_playbook_id,
               saved_playbook_check_id, latest_anchor_json, anchor_reconciliation_status,
               anchor_reconciliation_note, citations_json
        FROM review_suggestions
        WHERE review_run_id = ?
        ORDER BY created_at ASC
        """,
        (review_run_id,),
    ).fetchall()

    event_rows = connection.execute(
        """
        SELECT id, suggestion_id, action, actor_surface, status_after, note,
               client_message, applied_anchor_json, created_at
        FROM review_suggestion_events
        WHERE suggestion_id IN (
            SELECT id FROM review_suggestions WHERE review_run_id = ?
        )
        ORDER BY created_at DESC
        """,
        (review_run_id,),
    ).fetchall()
    event_map: dict[str, list[ReviewSuggestionEvent]] = {}
    for event_row in event_rows:
        event_map.setdefault(event_row["suggestion_id"], []).append(
            ReviewSuggestionEvent(
                id=event_row["id"],
                suggestion_id=event_row["suggestion_id"],
                action=event_row["action"],
                actor_surface=event_row["actor_surface"],
                status_after=SuggestionStatus(event_row["status_after"]),
                note=event_row["note"],
                created_at=event_row["created_at"],
                client_message=event_row["client_message"],
                applied_anchor=(
                    DocumentAnchor.model_validate(json.loads(event_row["applied_anchor_json"]))
                    if event_row["applied_anchor_json"]
                    and json.loads(event_row["applied_anchor_json"]) is not None
                    else None
                ),
            )
        )

    suggestions = [
        ReviewSuggestionRecord(
            id=suggestion_row["id"],
            review_run_id=suggestion_row["review_run_id"],
            anchor_id=suggestion_row["anchor_id"],
            title=suggestion_row["title"],
            issue_type=suggestion_row["issue_type"],
            severity=SeverityLevel(suggestion_row["severity"]),
            confidence=suggestion_row["confidence"],
            explanation=suggestion_row["explanation"],
            supporting_excerpt=suggestion_row["supporting_excerpt"],
            proposed_comment=suggestion_row["proposed_comment"],
            proposed_redline=(
                ProposedRedline.model_validate(
                    json.loads(suggestion_row["proposed_redline_json"])
                )
                if suggestion_row["proposed_redline_json"]
                and json.loads(suggestion_row["proposed_redline_json"]) is not None
                else None
            ),
            fallback_position_text=suggestion_row["fallback_position_text"],
            status=SuggestionStatus(suggestion_row["status"]),
            reviewer_note=suggestion_row["reviewer_note"],
            applied_at=suggestion_row["applied_at"],
            dismissed_at=suggestion_row["dismissed_at"],
            saved_to_playbook_at=suggestion_row["saved_to_playbook_at"],
            saved_playbook_note_id=suggestion_row["saved_playbook_note_id"],
            saved_playbook_id=suggestion_row["saved_playbook_id"],
            saved_playbook_check_id=suggestion_row["saved_playbook_check_id"],
            latest_anchor=(
                DocumentAnchor.model_validate(json.loads(suggestion_row["latest_anchor_json"]))
                if suggestion_row["latest_anchor_json"]
                and json.loads(suggestion_row["latest_anchor_json"]) is not None
                else None
            ),
            anchor_reconciliation_status=(
                AnchorReconciliationStatus(suggestion_row["anchor_reconciliation_status"])
                if suggestion_row["anchor_reconciliation_status"]
                else None
            ),
            anchor_reconciliation_note=suggestion_row["anchor_reconciliation_note"],
            citations=[
                CitationRecord.model_validate(entry)
                for entry in json.loads(suggestion_row["citations_json"])
            ],
            events=event_map.get(suggestion_row["id"], []),
        )
        for suggestion_row in suggestion_rows
    ]

    return ReviewRunRecord(
        id=row["id"],
        project_id=row["project_id"],
        document_version_id=row["document_version_id"],
        status=row["status"],
        review_type=ReviewType(row["review_type"]),
        represented_party=row["represented_party"],
        jurisdiction=row["jurisdiction"],
        audience=ReviewAudience(row["audience"]),
        summary=ReviewRunSummary.model_validate(json.loads(row["summary_json"])),
        suggestions=suggestions,
        created_at=row["created_at"],
        completed_at=row["completed_at"],
    )


def build_workspace_detail(connection: object, row: object) -> WorkspaceDetail:
    workspace_summary = build_workspace_summary(connection, row)
    documents = [
        DocumentRecord.model_validate(
            {
                **dict(document_row),
                "auto_renews": bool(document_row["auto_renews"]),
            }
        )
        for document_row in connection.execute(
            """
            SELECT id, name, doc_type, counterparty, effective_date, expiry_date,
                   renewal_notice_days, auto_renews, governing_law
            FROM documents
            WHERE workspace_id = ?
            ORDER BY created_at ASC
            """,
            (workspace_summary.id,),
        ).fetchall()
    ]

    issues = [
        IssueRecord.model_validate(
            {
                **dict(issue_row),
                "counterparties": json.loads(issue_row["counterparties_json"]),
                "key_dates": json.loads(issue_row["key_dates_json"]),
                "citations": json.loads(issue_row["citations_json"]),
            }
        )
        for issue_row in connection.execute(
            """
            SELECT id, workspace_id, document_id, title, issue_type, severity, status,
                   summary, reviewer_note, counterparties_json, key_dates_json, citations_json
            FROM issues
            WHERE workspace_id = ?
            ORDER BY created_at ASC
            """,
            (workspace_summary.id,),
        ).fetchall()
    ]

    detail = WorkspaceDetail(
        workspace=workspace_summary,
        documents=documents,
        issues=issues,
        highlights=[],
    )
    detail.highlights = build_highlights(detail)
    return detail


def build_workspace_id(name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    slug = slug or "workspace"
    return f"{slug}-{uuid4().hex[:6]}"


def sanitize_filename(filename: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "-", Path(filename).name)


def truncate_text(value: str, limit: int) -> str:
    normalized = " ".join(value.split())
    if len(normalized) <= limit:
        return normalized
    return f"{normalized[: max(limit - 3, 0)].rstrip()}..."


def normalize_clause_text(value: str) -> str:
    return " ".join(value.split()).strip()


def extract_clause_segments(value: str) -> list[str]:
    normalized = value.replace("\r", "\n")
    raw_segments = [
        normalize_clause_text(segment)
        for segment in re.split(r"\n{2,}", normalized)
        if normalize_clause_text(segment)
    ]
    segments = [segment[:650] for segment in raw_segments if len(segment) >= 40]
    if not segments:
        fallback = normalize_clause_text(value)
        return [fallback[:650]] if fallback else []
    return segments[:6]


def derive_library_item_title(text: str, section_heading: str) -> str:
    preview = truncate_text(text, 72)
    if section_heading and section_heading.lower() != "document text":
        return f"{section_heading}: {preview}"
    return preview


def validate_upload(filename: str) -> None:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise ValueError(f"Unsupported file type for {filename}. Only PDF and DOCX are accepted.")


def now_timestamp() -> str:
    return datetime.now(tz=UTC).isoformat()


def build_review_run_summary_markdown(review_run: ReviewRunRecord) -> str:
    lines = [
        f"# Review Summary: {review_run.review_type.title()}",
        "",
        f"- Review run: {review_run.id}",
        f"- Project: {review_run.project_id}",
        f"- Document version: {review_run.document_version_id}",
        f"- Represented party: {review_run.represented_party}",
        f"- Jurisdiction: {review_run.jurisdiction}",
        f"- Audience: {review_run.audience}",
        f"- Exported at: {now_timestamp()}",
        "",
        "## Counts",
        f"- Total: {review_run.summary.total if review_run.summary else 0}",
        f"- High: {review_run.summary.high if review_run.summary else 0}",
        f"- Medium: {review_run.summary.medium if review_run.summary else 0}",
        f"- Low: {review_run.summary.low if review_run.summary else 0}",
        "",
        "## Suggestions",
    ]

    if not review_run.suggestions:
        lines.append("- No suggestions were generated.")
        return "\n".join(lines)

    for suggestion in review_run.suggestions:
        lines.extend(
            [
                f"### {suggestion.title}",
                f"- Type: {suggestion.issue_type}",
                f"- Severity: {suggestion.severity}",
                f"- Status: {suggestion.status}",
                f"- Confidence: {round(suggestion.confidence * 100)}%",
                f"- Excerpt: {suggestion.supporting_excerpt}",
                f"- Explanation: {suggestion.explanation}",
                "",
            ]
        )

    return "\n".join(lines)


def resolve_target_playbook(playbook_id: str | None) -> PlaybookRecord | None:
    playbooks = load_playbooks()

    if playbook_id and playbook_id.strip():
        requested = playbook_id.strip()
        for playbook in playbooks:
            if playbook.name == requested:
                return playbook

    return playbooks[0] if playbooks else None


def resolve_target_playbook_check_id(
    target_playbook: PlaybookRecord | None,
    requested_check_id: str | None,
    issue_type: str,
) -> str | None:
    if target_playbook is None:
        return None

    if requested_check_id and requested_check_id.strip():
        requested = requested_check_id.strip()
        for check in target_playbook.checks:
            if check.id == requested:
                return check.id

    normalized_issue_type = issue_type.strip().lower()
    for check in target_playbook.checks:
        if check.id.lower() == normalized_issue_type:
            return check.id

    return None


def build_saved_to_playbook_message(playbook_id: str, playbook_check_id: str | None) -> str:
    if playbook_check_id:
        return f"Saved to playbook {playbook_id} / {playbook_check_id}."

    return f"Saved to playbook {playbook_id}."


def resolve_saved_playbook_note(
    explicit_note: str | None,
    reviewer_note: str | None,
    supporting_excerpt: str,
) -> str:
    for candidate in (explicit_note, reviewer_note):
        if candidate and candidate.strip():
            return candidate.strip()

    return supporting_excerpt.strip()


def resolve_original_anchor_quote(supporting_excerpt: str, citations_json: str) -> str:
    citations = json.loads(citations_json)
    if citations:
        first_quote = citations[0].get("quote")
        if first_quote:
            return str(first_quote)

    return supporting_excerpt


def reconcile_anchor_state(
    mode: str,
    original_excerpt: str,
    applied_anchor: DocumentAnchor | None,
) -> tuple[AnchorReconciliationStatus, str]:
    if applied_anchor is None or not applied_anchor.quote.strip():
        return (
            AnchorReconciliationStatus.UNKNOWN,
            "No post-apply anchor snapshot was returned from Word.",
        )

    original_tokens = tokenize_anchor_text(original_excerpt)
    applied_tokens = tokenize_anchor_text(applied_anchor.quote)

    if not original_tokens or not applied_tokens:
        return (
            AnchorReconciliationStatus.UNKNOWN,
            "Anchor reconciliation could not compare empty text.",
        )

    overlap_ratio = len(original_tokens & applied_tokens) / max(len(original_tokens), 1)

    if mode == "comment":
        if overlap_ratio >= 0.8:
            return (
                AnchorReconciliationStatus.STABLE,
                "The post-apply anchor still closely matches the original selection.",
            )

        return (
            AnchorReconciliationStatus.DRIFTED,
            "The post-apply selection differs materially from the original clause after comment application.",
        )

    if overlap_ratio >= 0.35:
        return (
            AnchorReconciliationStatus.UPDATED,
            "The returned selection changed as expected after the redline, but still overlaps the original clause.",
        )

    return (
        AnchorReconciliationStatus.DRIFTED,
        "The returned post-redline selection has very little overlap with the original clause. Verify the anchor.",
    )


def tokenize_anchor_text(value: str) -> set[str]:
    tokens = re.findall(r"[a-z0-9]+", value.lower())
    return set(tokens)
